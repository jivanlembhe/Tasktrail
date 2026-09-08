#!/usr/bin/env python3
"""
TaskTrail — task planner, to-do list and status tracker (Python / PySide6 edition)

Data-compatible with the Electron TaskTrail: same JSON format, same file
(%APPDATA%\\TaskTrail\\flowboard_data.json on Windows), same backups.

Run:     python tasktrail.py
Build:   see README_PYTHON.md (PyInstaller one-liner)
"""
import copy
import datetime as dt
import json
import os
import shutil
import sys
import time
import uuid

from PySide6.QtCore import Qt, QTimer, QSize, QRect, QPoint, Signal, QDate
from PySide6.QtGui import (QAction, QColor, QFont, QIcon, QPainter, QPixmap, QBrush, QPen,
                           QFontDatabase, QCursor, QShortcut, QKeySequence)
from PySide6.QtWidgets import (QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout, QGridLayout,
                               QLabel, QPushButton, QFrame, QStackedWidget, QListWidget, QListWidgetItem,
                               QAbstractItemView, QScrollArea, QLineEdit, QComboBox, QTextEdit, QDialog,
                               QDialogButtonBox, QCheckBox, QDateEdit, QColorDialog, QFileDialog, QMessageBox,
                               QSystemTrayIcon, QMenu, QSizePolicy, QToolButton, QDockWidget, QSpinBox,
                               QFontComboBox, QSlider, QInputDialog, QSplitter)

APP_NAME = "TaskTrail"
VERSION = "2.0.1-py"
MONTHS = ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec']
MONTHS_LONG = ['January', 'February', 'March', 'April', 'May', 'June', 'July', 'August',
               'September', 'October', 'November', 'December']
DEFAULT_COLS = [
    {"id": "todo", "label": "To Do", "color": "#6c8aff", "icon": "📋"},
    {"id": "today", "label": "Today's", "color": "#ffb347", "icon": "☀️"},
    {"id": "wip", "label": "WIP", "color": "#b48aff", "icon": "⚡"},
    {"id": "done", "label": "Done", "color": "#3ddbbf", "icon": "✅"},
]
BUILTIN = {"todo", "today", "wip", "done"}
LEGACY_YEAR = 2026
LABEL_PALETTE = ['#6c8aff', '#ffb347', '#b48aff', '#3ddbbf', '#ff6584', '#26de81',
                 '#fd9644', '#a55eea', '#2bcbba', '#eb3b5a', '#f7b731', '#778ca3']
PRI_LABEL = {"high": "🔴 High", "med": "🟡 Med", "low": "🟢 Low"}
PRI_COLOR = {"high": "#ff6584", "med": "#ffb347", "low": "#3ddbbf"}

THEMES = {
    "dark": dict(bg="#07080d", s1="#0e1018", s2="#13161f", s3="#191d29", s4="#1f2333", border="#252a3a",
                 border2="#2e3448", text="#eef0f8", text2="#9aa0bc", text3="#5c6382", accent="#6c8aff"),
    "light": dict(bg="#f0f2f9", s1="#ffffff", s2="#ffffff", s3="#f4f5fb", s4="#eaecf5", border="#d8dbee",
                  border2="#c4c8e0", text="#1a1d2e", text2="#4a5070", text3="#8890b0", accent="#4a6bef"),
}
PRESETS = [
    ("Default dark", "dark", {}),
    ("Default light", "light", {}),
    ("Midnight", "dark", dict(bg="#0a1020", surface="#121a30", text="#e8ecff", accent="#7c9cff")),
    ("Forest", "dark", dict(bg="#0b1410", surface="#12211a", text="#e6f2ec", accent="#3ddbbf")),
    ("Ember", "dark", dict(bg="#16100d", surface="#241a16", text="#f7ece6", accent="#ff8c5a")),
    ("Paper", "light", dict(bg="#f5f2ea", surface="#fffdf8", text="#25231f", accent="#b5622e")),
]


# ═══════════════════════════════════════════════════════════════════════════
#  PATHS & SETTINGS
# ═══════════════════════════════════════════════════════════════════════════
def user_data_dir():
    if sys.platform.startswith("win"):
        base = os.environ.get("APPDATA") or os.path.expanduser("~")
    elif sys.platform == "darwin":
        base = os.path.expanduser("~/Library/Application Support")
    else:
        base = os.environ.get("XDG_CONFIG_HOME") or os.path.expanduser("~/.config")
    d = os.path.join(base, APP_NAME)
    os.makedirs(d, exist_ok=True)
    return d


USER_DATA = user_data_dir()
DB_FILE = os.path.join(USER_DATA, "flowboard_data.json")
SETTINGS_FILE = os.path.join(USER_DATA, "settings.json")
DEFAULT_BACKUP_DIR = os.path.join(USER_DATA, "backups")


def read_settings():
    try:
        with open(SETTINGS_FILE, encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return {}


def write_settings(s):
    with open(SETTINGS_FILE, "w", encoding="utf-8") as f:
        json.dump(s, f, indent=2)


def get_backup_dir():
    d = read_settings().get("backupDir") or DEFAULT_BACKUP_DIR
    try:
        os.makedirs(d, exist_ok=True)
        return d
    except Exception:
        os.makedirs(DEFAULT_BACKUP_DIR, exist_ok=True)
        return DEFAULT_BACKUP_DIR


def detect_onedrive():
    for p in (os.environ.get("OneDriveCommercial"), os.environ.get("OneDriveConsumer"),
              os.environ.get("OneDrive"), os.path.join(os.path.expanduser("~"), "OneDrive")):
        if p and os.path.isdir(p):
            return p
    return None


def rgba(hex_color, a):
    """Qt stylesheets don't accept #RRGGBBAA — build rgba() instead."""
    h = hex_color.lstrip('#')
    if len(h) != 6:
        return hex_color
    r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    return f"rgba({r},{g},{b},{a})"


def uid():
    return uuid.uuid4().hex[:12]


def now_iso():
    return dt.datetime.now().isoformat(timespec="seconds")


def fmt_date(s):
    try:
        return dt.date.fromisoformat(s).strftime("%b %d").replace(" 0", " ")
    except Exception:
        return s or ""


def fmt_dt(s):
    try:
        return dt.datetime.fromisoformat(s.replace("Z", "")).strftime("%b %d, %H:%M")
    except Exception:
        return ""


def is_overdue(d, done):
    if not d or done:
        return False
    try:
        return dt.date.fromisoformat(d) < dt.date.today()
    except Exception:
        return False


# ═══════════════════════════════════════════════════════════════════════════
#  STORE — same JSON schema as the Electron app
# ═══════════════════════════════════════════════════════════════════════════
class Store:
    def __init__(self):
        self.db = {}
        self.year = dt.date.today().year
        self.month = dt.date.today().month - 1
        self.load()
        if not os.path.exists(DB_FILE):
            self.save()                 # make sure the file exists from the first run so backups work

    # ---- load / save ----
    def load(self):
        raw = None
        if os.path.exists(DB_FILE):
            try:
                with open(DB_FILE, encoding="utf-8") as f:
                    raw = json.load(f)
            except Exception:
                raw = None
        self.db = self.normalize(raw)
        self.ensure_year(self.year)

    @staticmethod
    def normalize(raw):
        d = raw if isinstance(raw, dict) else {}
        if not isinstance(d.get("years"), dict):
            d["years"] = {}
            if d.get("kanban") or d.get("checklist"):
                d["years"][str(LEGACY_YEAR)] = {"kanban": d.get("kanban") or {}, "checklist": d.get("checklist") or {}}
        d.pop("kanban", None)
        d.pop("checklist", None)
        if not isinstance(d.get("columns"), list) or not d["columns"]:
            d["columns"] = copy.deepcopy(DEFAULT_COLS)
        for dc in DEFAULT_COLS:
            if not any(c["id"] == dc["id"] for c in d["columns"]):
                d["columns"].append(dict(dc))
        d.setdefault("labels", [])
        d.setdefault("migrations", [])
        for m in d["migrations"]:
            m.setdefault("fromYear", LEGACY_YEAR)
            m.setdefault("toYear", LEGACY_YEAR)
        if not isinstance(d.get("meta"), dict):
            d["meta"] = {}
        return d

    def ensure_year(self, y):
        Y = self.db["years"].setdefault(str(y), {"kanban": {}, "checklist": {}})
        Y.setdefault("kanban", {})
        Y.setdefault("checklist", {})
        for mi in range(12):
            k = Y["kanban"].setdefault(str(mi), {})
            for c in self.db["columns"]:
                k.setdefault(c["id"], [])
            Y["checklist"].setdefault(str(mi), [])
        return Y

    def save(self):
        self.db["meta"]["updatedAt"] = int(time.time() * 1000)
        tmp = DB_FILE + ".tmp"
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump(self.db, f)
        os.replace(tmp, DB_FILE)

    # ---- accessors ----
    @property
    def columns(self):
        return self.db["columns"]

    @property
    def labels(self):
        return self.db["labels"]

    def col(self, cid):
        return next((c for c in self.columns if c["id"] == cid), None)

    def kanban(self, mi=None, year=None):
        Y = self.ensure_year(year or self.year)
        return Y["kanban"][str(self.month if mi is None else mi)]

    def checklist(self, mi=None, year=None):
        Y = self.ensure_year(year or self.year)
        return Y["checklist"][str(self.month if mi is None else mi)]

    def years(self):
        return sorted(int(y) for y in self.db["years"])

    def find_card(self, cid, col_id=None):
        kd = self.kanban()
        for c in self.columns:
            if col_id and c["id"] != col_id:
                continue
            for card in kd.get(c["id"], []):
                if card["id"] == cid:
                    return card, c["id"]
        return None, None

    @staticmethod
    def log(card, text):
        card.setdefault("activity", []).append({"at": now_iso(), "text": text})
        del card["activity"][:-60]

    def move_card(self, cid, from_col, to_col, before_id="TOP"):
        src = self.kanban()[from_col]
        idx = next((i for i, c in enumerate(src) if c["id"] == cid), -1)
        if idx < 0:
            return None
        card = src.pop(idx)
        dst = self.kanban().setdefault(to_col, [])
        if before_id == "TOP":
            at = 0
        elif before_id is None:
            at = len(dst)
        else:
            at = next((i for i, c in enumerate(dst) if c["id"] == before_id), len(dst))
        dst.insert(at, card)
        if from_col != to_col:
            card["done"] = to_col == "done"
            if to_col == "done":
                card["completedAt"] = now_iso()
            self.log(card, f"Moved from {self.col(from_col)['label']} to {self.col(to_col)['label']}")
        return card

    # ---- rollover ----
    def incomplete_col_ids(self):
        return [c["id"] for c in self.columns if c["id"] != "done"]

    def incomplete_cards(self, mi, year):
        kd = self.kanban(mi, year)
        return [c for cid in self.incomplete_col_ids() for c in kd.get(cid, []) if not c.get("done")]

    def incomplete_cl(self, mi, year):
        return [t for g in self.checklist(mi, year) for t in g["tasks"] if not t.get("done")]

    @staticmethod
    def month_past(mi, year):
        last = dt.date(year + (1 if mi == 11 else 0), 1 if mi == 11 else mi + 2, 1) - dt.timedelta(days=1)
        return dt.date.today() > last

    def past_periods(self):
        today = dt.date.today()
        out = []
        for y in self.years():
            if y > today.year:
                continue
            last_m = 11 if y < today.year else today.month - 2
            for mi in range(0, last_m + 1):
                if self.month_past(mi, y):
                    out.append((y, mi))
        return out

    def migrate_month(self, mi, year, auto=False):
        ty, tm = (year + 1, 0) if mi >= 11 else (year, mi + 1)
        self.ensure_year(year); self.ensure_year(ty)
        src, tgt = self.kanban(mi, year), self.kanban(tm, ty)
        label = f"{MONTHS[mi]} {year}"
        moved_k = moved_c = 0
        for cid in self.incomplete_col_ids():
            arr = src.get(cid, [])
            for c in [c for c in arr if not c.get("done")]:
                m = copy.deepcopy(c)
                m.update(id=uid(), migratedFrom=label, migratedAt=now_iso(), originalId=c["id"])
                self.log(m, f"Rolled over from {label}")
                tgt.setdefault("todo" if cid == "today" else cid, []).insert(0, m)
                moved_k += 1
            src[cid] = [c for c in arr if c.get("done")]
        for g in self.checklist(mi, year):
            inc = [t for t in g["tasks"] if not t.get("done")]
            if not inc:
                continue
            tg = next((x for x in self.checklist(tm, ty) if x["name"] == g["name"]), None)
            if not tg:
                tg = {"id": uid(), "name": g["name"], "tasks": []}
                self.checklist(tm, ty).append(tg)
            for t in inc:
                m = copy.deepcopy(t); m.update(id=uid(), migratedFrom=label, migratedAt=now_iso(), originalId=t["id"])
                tg["tasks"].insert(0, m); moved_c += 1
            g["tasks"] = [t for t in g["tasks"] if t.get("done")]
        self.db["migrations"].append({"fromYear": year, "from": mi, "toYear": ty, "to": tm, "kanban": moved_k,
                                      "checklist": moved_c, "at": now_iso(), "auto": auto})
        self.save()
        return moved_k, moved_c

    def auto_migrate(self):
        total = 0
        for y, mi in self.past_periods():
            if not self.incomplete_cards(mi, y) and not self.incomplete_cl(mi, y):
                continue
            if any(m.get("fromYear") == y and m["from"] == mi and m.get("auto") for m in self.db["migrations"]):
                continue
            k, c = self.migrate_month(mi, y, auto=True)
            total += k + c
        return total

    def pending_rollover(self):
        return any(self.incomplete_cards(mi, y) or self.incomplete_cl(mi, y) for y, mi in self.past_periods())

    # ---- backups ----
    def backup_now(self):
        if not os.path.exists(DB_FILE):
            return None
        d = get_backup_dir()
        dest = os.path.join(d, "flowboard_backup_" + dt.datetime.now().strftime("%Y-%m-%dT%H-%M-%S-%f")[:-3] + ".json")
        shutil.copyfile(DB_FILE, dest)
        files = sorted([f for f in os.listdir(d) if f.startswith("flowboard_backup_") and f.endswith(".json")],
                       key=lambda f: os.path.getmtime(os.path.join(d, f)), reverse=True)
        for f in files[30:]:
            try:
                os.remove(os.path.join(d, f))
            except Exception:
                pass
        return dest

    @staticmethod
    def backup_list():
        d = get_backup_dir()
        out = []
        for f in os.listdir(d):
            if f.endswith(".json"):
                p = os.path.join(d, f)
                out.append((f, os.path.getsize(p), os.path.getmtime(p)))
        return sorted(out, key=lambda x: x[2], reverse=True)

    def restore_file(self, path):
        with open(path, encoding="utf-8") as f:
            parsed = json.load(f)
        if not isinstance(parsed, dict) or not (parsed.get("kanban") or parsed.get("years")):
            raise ValueError("Not a TaskTrail / FlowBoard backup file.")
        data = open(path, "rb").read()          # read first: the safety backup must never clobber the source
        self.backup_now()
        with open(DB_FILE, "wb") as f:
            f.write(data)
        self.load()
        self.save()


# ═══════════════════════════════════════════════════════════════════════════
#  THEME / STYLESHEET
# ═══════════════════════════════════════════════════════════════════════════
def build_qss(p, font_family, font_size):
    return f"""
    QWidget {{ background:{p['bg']}; color:{p['text']}; font-family:"{font_family}"; font-size:{font_size}px; }}
    QMainWindow, QDialog {{ background:{p['bg']}; }}
    QFrame#sidebar {{ background:{p['s1']}; border-right:1px solid {p['border']}; }}
    QFrame#topbar {{ background:{p['s1']}; border-bottom:1px solid {p['border']}; }}
    QLabel#brand {{ font-size:{font_size+8}px; font-weight:800; color:{p['accent']}; }}
    QLabel#brandSub, QLabel#navSection {{ color:{p['text3']}; font-size:{font_size-4}px; letter-spacing:1px; }}
    QLabel#pageTitle {{ font-size:{font_size+5}px; font-weight:700; }}
    QLabel#pageMonth {{ color:{p['accent']}; background:{rgba(p['accent'], 0.14)}; border-radius:10px; font-weight:600; }}
    QPushButton#nav {{ text-align:left; padding:8px 18px; border:none; border-left:3px solid transparent; color:{p['text2']}; background:transparent; font-weight:500; }}
    QPushButton#nav:hover {{ color:{p['text']}; background:{p['s2']}; }}
    QPushButton#nav:checked {{ color:{p['accent']}; background:{rgba(p['accent'], 0.12)}; border-left:3px solid {p['accent']}; }}
    QPushButton#month {{ padding:4px 2px; border:1px solid {p['border']}; border-radius:6px; color:{p['text3']}; background:transparent; font-size:{font_size-3}px; font-weight:700; }}
    QPushButton#month:hover {{ border-color:{p['accent']}; color:{p['accent']}; }}
    QPushButton#month:checked {{ background:{p['accent']}; border-color:{p['accent']}; color:#ffffff; }}
    QPushButton {{ padding:6px 14px; border-radius:14px; border:1.5px solid {p['border']}; background:transparent; color:{p['text2']}; font-weight:700; }}
    QPushButton:hover {{ color:{p['text']}; border-color:{p['text2']}; }}
    QPushButton#primary {{ background:{p['accent']}; border-color:{p['accent']}; color:#ffffff; }}
    QPushButton#primary:hover {{ background:{p['accent']}dd; }}
    QPushButton#danger {{ color:#ff6584; border-color:#ff6584; }}
    QPushButton#ghost {{ border:none; color:{p['text3']}; padding:2px 6px; }}
    QPushButton#ghost:hover {{ color:{p['text']}; background:{p['s3']}; border-radius:6px; }}
    QPushButton#addk {{ border:1.5px dashed {p['border']}; border-radius:12px; color:{p['text3']}; font-weight:500; text-align:left; padding:8px 14px; }}
    QPushButton#addk:hover {{ border-color:{p['accent']}; color:{p['accent']}; }}
    QLineEdit, QTextEdit, QComboBox, QDateEdit, QSpinBox, QFontComboBox {{ background:{p['s3']}; border:1.5px solid {p['border']}; border-radius:8px; padding:6px 10px; color:{p['text']}; selection-background-color:{p['accent']}; }}
    QLineEdit:focus, QTextEdit:focus, QComboBox:focus, QDateEdit:focus {{ border-color:{p['accent']}; }}
    QComboBox QAbstractItemView {{ background:{p['s2']}; color:{p['text']}; selection-background-color:{p['accent']}; border:1px solid {p['border']}; }}
    QFrame#panel, QFrame#kpi, QFrame#card, QFrame#colhead, QFrame#group {{ background:{p['s2']}; border:1px solid {p['border']}; border-radius:12px; }}
    QFrame#card:hover {{ border-color:{p['border2']}; }}
    QFrame#detail {{ background:{p['s1']}; border-left:1px solid {p['border']}; }}
    QListWidget {{ background:transparent; border:none; outline:0; }}
    QListWidget::item {{ background:transparent; border:none; padding:0; margin:0 0 6px 0; }}
    QListWidget::item:selected {{ background:transparent; }}
    QScrollArea {{ border:none; background:transparent; }}
    QScrollBar:vertical {{ width:10px; background:transparent; }} QScrollBar:horizontal {{ height:10px; background:transparent; }}
    QScrollBar::handle {{ background:{p['border2']}; border-radius:5px; min-height:24px; min-width:24px; }}
    QScrollBar::handle:hover {{ background:{p['text3']}; }}
    QScrollBar::add-line, QScrollBar::sub-line {{ height:0; width:0; }}
    QLabel, QCheckBox {{ background:transparent; }}
    QLabel#muted {{ color:{p['text3']}; }} QLabel#muted2 {{ color:{p['text2']}; }}
    QLabel#kpiVal {{ font-size:{font_size+16}px; font-weight:700; }}
    QLabel#kpiLabel {{ color:{p['text2']}; font-size:{font_size-3}px; letter-spacing:1px; }}
    QLabel#sectitle {{ font-weight:700; font-size:{font_size+1}px; }}
    QLabel#badge {{ border-radius:9px; padding:1px 7px; font-size:{font_size-4}px; font-weight:700; }}
    QCheckBox {{ spacing:8px; }} QCheckBox::indicator {{ width:16px; height:16px; border-radius:4px; border:2px solid {p['border2']}; background:transparent; }}
    QCheckBox::indicator:checked {{ background:#3ddbbf; border-color:#3ddbbf; }}
    QMenu {{ background:{p['s2']}; border:1px solid {p['border']}; }} QMenu::item:selected {{ background:{p['accent']}; color:#fff; }}
    QToolTip {{ background:{p['s2']}; color:{p['text']}; border:1px solid {p['border']}; }}
    QDockWidget::title {{ background:{p['s1']}; padding:6px; }}
    """


# ═══════════════════════════════════════════════════════════════════════════
#  SMALL WIDGETS
# ═══════════════════════════════════════════════════════════════════════════
def badge(text, color, bg=None):
    l = QLabel(text)
    l.setObjectName("badge")
    l.setStyleSheet(f"QLabel#badge{{color:{color};background:{bg or rgba(color, 0.15)};}}")
    l.setMargin(2)
    return l


def hline():
    f = QFrame(); f.setFrameShape(QFrame.HLine); f.setStyleSheet("color: palette(mid);")
    return f


class BarChart(QWidget):
    """Tasks per month (current year)."""
    def __init__(self, get_values, accent):
        super().__init__(); self.get_values = get_values; self.accent = accent
        self.setMinimumHeight(110)

    def paintEvent(self, e):
        vals = self.get_values()
        mx = max(1, max(vals))
        p = QPainter(self); p.setRenderHint(QPainter.Antialiasing)
        w = self.width(); h = self.height() - 20
        bw = (w - 11 * 6) / 12
        for i, v in enumerate(vals):
            bh = max(4, int(v / mx * (h - 8)))
            x = int(i * (bw + 6))
            p.setPen(Qt.NoPen); p.setBrush(QColor(self.accent))
            p.drawRoundedRect(x, h - bh, int(bw), bh, 3, 3)
            p.setPen(QColor("#5c6382")); p.drawText(QRect(x, h + 2, int(bw), 16), Qt.AlignCenter, MONTHS[i][0])
        p.end()


class CardWidget(QFrame):
    """A task card on the board."""
    clicked = Signal(str, str)      # cid, col
    edit = Signal(str, str)
    delete = Signal(str, str)
    toggle = Signal(str, str)

    def __init__(self, store, card, col):
        super().__init__(); self.setObjectName("card"); self.card = card; self.col_id = col["id"]
        self.setCursor(QCursor(Qt.PointingHandCursor))
        self.setStyleSheet(f"QFrame#card{{border-left:3px solid {col['color']};}}")
        v = QVBoxLayout(self); v.setContentsMargins(12, 10, 10, 10); v.setSpacing(6)
        labels = [l for l in store.labels if l["id"] in card.get("labels", [])]
        if labels:
            lr = QHBoxLayout(); lr.setSpacing(4)
            for l in labels:
                lr.addWidget(badge(l["name"], "#ffffff", l["color"]))
            lr.addStretch(); v.addLayout(lr)
        top = QHBoxLayout(); top.setSpacing(8)
        done = card.get("done") or self.col_id == "done"
        cb = QCheckBox(); cb.setChecked(done); cb.setToolTip("Complete / reopen")
        cb.clicked.connect(lambda: self.toggle.emit(card["id"], self.col_id))
        title = QLabel(card["title"]); title.setWordWrap(True); title.setStyleSheet("font-weight:500;" + ("text-decoration:line-through;color:#5c6382;" if done else ""))
        title.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Preferred)
        eb = QPushButton("✎"); eb.setObjectName("ghost"); eb.setToolTip("Edit"); eb.clicked.connect(lambda: self.edit.emit(card["id"], self.col_id))
        db_ = QPushButton("✕"); db_.setObjectName("ghost"); db_.setToolTip("Delete"); db_.clicked.connect(lambda: self.delete.emit(card["id"], self.col_id))
        for w in (cb, title, eb, db_):
            top.addWidget(w)
        v.addLayout(top)
        meta = QHBoxLayout(); meta.setSpacing(5)
        meta.addWidget(badge(PRI_LABEL[card.get("priority", "med")], PRI_COLOR[card.get("priority", "med")]))
        if card.get("dueDate"):
            od = is_overdue(card["dueDate"], done)
            meta.addWidget(badge(("⚠ " if od else "📅 ") + fmt_date(card["dueDate"]), "#ff6584" if od else "#8890b0"))
        if card.get("migratedFrom"):
            meta.addWidget(badge("↪ " + card["migratedFrom"], "#8890b0"))
        meta.addStretch(); v.addLayout(meta)
        if card.get("desc"):
            d = QLabel(card["desc"][:160] + ("…" if len(card["desc"]) > 160 else "")); d.setObjectName("muted2"); d.setWordWrap(True); v.addWidget(d)
        subs = card.get("subs", [])
        if subs:
            sd = sum(1 for s in subs if s.get("done"))
            v.addWidget(QLabel(f"☑ {sd}/{len(subs)} sub-tasks", objectName="muted"))
        if card.get("comments"):
            v.addWidget(QLabel(f"💬 {len(card['comments'])}", objectName="muted"))

    def mouseReleaseEvent(self, e):
        if e.button() == Qt.LeftButton:
            self.clicked.emit(self.card["id"], self.col_id)
        super().mouseReleaseEvent(e)


class CardList(QListWidget):
    """Column list with drag & drop between/within columns. The drop updates the model
    and asks the board to re-render, so item widgets are always rebuilt."""
    dropped = Signal(str, str, str, object)   # cid, from_col, to_col, before_id(None=end)

    def __init__(self, col_id):
        super().__init__(); self.col_id = col_id
        self.setDragDropMode(QAbstractItemView.DragDrop); self.setDefaultDropAction(Qt.MoveAction)
        self.setAcceptDrops(True); self.setDragEnabled(True); self.setSelectionMode(QAbstractItemView.SingleSelection)
        self.setVerticalScrollMode(QAbstractItemView.ScrollPerPixel); self.setSpacing(0)
        self.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        self.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)
        self.setUniformItemSizes(False)

    def resizeEvent(self, e):
        super().resizeEvent(e)
        self.fit_items()

    def fit_items(self):
        """Size every card to the column width so wrapped text and borders are never clipped."""
        w = self.viewport().width() - 2
        if w < 60:
            return
        for i in range(self.count()):
            it = self.item(i)
            wd = self.itemWidget(it)
            if wd is None:
                continue
            wd.setFixedWidth(w)
            h = wd.heightForWidth(w) if wd.layout() and wd.layout().hasHeightForWidth() else wd.sizeHint().height()
            it.setSizeHint(QSize(w, max(h, wd.sizeHint().height()) + 6))

    def dropEvent(self, e):
        src = e.source()
        if not isinstance(src, CardList) or not src.currentItem():
            e.ignore(); return
        cid = src.currentItem().data(Qt.UserRole)
        pos = e.position().toPoint() if hasattr(e, "position") else e.pos()
        item = self.itemAt(pos)
        before = None
        if item is not None:
            r = self.visualItemRect(item)
            if pos.y() < r.center().y():
                before = item.data(Qt.UserRole)
            else:
                nxt = self.item(self.row(item) + 1)
                before = nxt.data(Qt.UserRole) if nxt else None
        e.setDropAction(Qt.IgnoreAction); e.accept()
        self.dropped.emit(cid, src.col_id, self.col_id, before)


# ═══════════════════════════════════════════════════════════════════════════
#  DIALOGS
# ═══════════════════════════════════════════════════════════════════════════
class TaskDialog(QDialog):
    def __init__(self, parent, store, card=None, col_id="todo"):
        super().__init__(parent); self.store = store; self.card = card
        self.setWindowTitle("Edit Task" if card else "New Task"); self.setMinimumWidth(480)
        v = QVBoxLayout(self); v.setSpacing(10)
        self.title = QLineEdit(card["title"] if card else ""); self.title.setPlaceholderText("What needs to be done?")
        self.desc = QTextEdit(card.get("desc", "") if card else ""); self.desc.setPlaceholderText("Details or notes…"); self.desc.setFixedHeight(70)
        self.col = QComboBox()
        for c in store.columns:
            self.col.addItem(f"{c.get('icon', '')} {c['label']}", c["id"])
        self.col.setCurrentIndex(max(0, self.col.findData(col_id)))
        self.pri = QComboBox()
        for k, lab in PRI_LABEL.items():
            self.pri.addItem(lab, k)
        self.pri.setCurrentIndex(self.pri.findData(card.get("priority", "med") if card else "med"))
        self.has_date = QCheckBox("Due date"); self.date = QDateEdit(QDate.currentDate()); self.date.setCalendarPopup(True); self.date.setDisplayFormat("yyyy-MM-dd")
        if card and card.get("dueDate"):
            self.has_date.setChecked(True); self.date.setDate(QDate.fromString(card["dueDate"], "yyyy-MM-dd"))
        self.date.setEnabled(self.has_date.isChecked()); self.has_date.toggled.connect(self.date.setEnabled)
        v.addWidget(QLabel("Title *")); v.addWidget(self.title)
        v.addWidget(QLabel("Description")); v.addWidget(self.desc)
        g = QGridLayout(); g.addWidget(QLabel("Column"), 0, 0); g.addWidget(QLabel("Priority"), 0, 1)
        g.addWidget(self.col, 1, 0); g.addWidget(self.pri, 1, 1); v.addLayout(g)
        dr = QHBoxLayout(); dr.addWidget(self.has_date); dr.addWidget(self.date); dr.addStretch(); v.addLayout(dr)
        # labels
        lr = QHBoxLayout(); lr.addWidget(QLabel("Labels")); mb = QPushButton("Manage"); mb.clicked.connect(self.manage_labels); lr.addWidget(mb); lr.addStretch(); v.addLayout(lr)
        self.label_box = QWidget(); self.label_layout = QHBoxLayout(self.label_box); self.label_layout.setContentsMargins(0, 0, 0, 0)
        self.sel_labels = set(card.get("labels", [])) if card else set(); self.rebuild_labels(); v.addWidget(self.label_box)
        # subtasks
        v.addWidget(QLabel("Sub-tasks"))
        self.subs = [dict(s) for s in (card.get("subs", []) if card else [])]
        self.sub_list = QListWidget(); self.sub_list.setFixedHeight(110); self.rebuild_subs(); v.addWidget(self.sub_list)
        sr = QHBoxLayout(); self.sub_inp = QLineEdit(); self.sub_inp.setPlaceholderText("Add sub-task… (Enter)")
        self.sub_inp.returnPressed.connect(self.add_sub); ab = QPushButton("+"); ab.clicked.connect(self.add_sub); rb = QPushButton("Remove selected"); rb.clicked.connect(self.remove_sub)
        sr.addWidget(self.sub_inp); sr.addWidget(ab); sr.addWidget(rb); v.addLayout(sr)
        bb = QDialogButtonBox(QDialogButtonBox.Save | QDialogButtonBox.Cancel)
        bb.button(QDialogButtonBox.Save).setObjectName("primary"); bb.button(QDialogButtonBox.Save).setText("Save changes" if card else "Save task")
        bb.accepted.connect(self.accept); bb.rejected.connect(self.reject); v.addWidget(bb)
        QShortcut(QKeySequence("Ctrl+Return"), self, activated=self.accept)

    def rebuild_labels(self):
        while self.label_layout.count():
            w = self.label_layout.takeAt(0).widget()
            if w: w.deleteLater()
        if not self.store.labels:
            self.label_layout.addWidget(QLabel("No labels yet", objectName="muted"))
        for l in self.store.labels:
            b = QPushButton(l["name"]); b.setCheckable(True); b.setChecked(l["id"] in self.sel_labels)
            b.setStyleSheet(f"QPushButton{{border-color:{l['color']};color:{l['color']};}}QPushButton:checked{{background:{l['color']};color:#fff;}}")
            b.toggled.connect(lambda on, lid=l["id"]: self.sel_labels.add(lid) if on else self.sel_labels.discard(lid))
            self.label_layout.addWidget(b)
        self.label_layout.addStretch()

    def manage_labels(self):
        LabelDialog(self, self.store).exec(); self.rebuild_labels()

    def rebuild_subs(self):
        self.sub_list.clear()
        for s in self.subs:
            self.sub_list.addItem(("✓ " if s.get("done") else "○ ") + s["text"])

    def add_sub(self):
        t = self.sub_inp.text().strip()
        if t:
            self.subs.append({"id": uid(), "text": t, "done": False}); self.sub_inp.clear(); self.rebuild_subs()

    def remove_sub(self):
        r = self.sub_list.currentRow()
        if r >= 0:
            del self.subs[r]; self.rebuild_subs()

    def accept(self):
        if not self.title.text().strip():
            self.title.setFocus(); return
        super().accept()

    def values(self):
        return dict(title=self.title.text().strip(), desc=self.desc.toPlainText().strip(), col=self.col.currentData(),
                    priority=self.pri.currentData(), dueDate=self.date.date().toString("yyyy-MM-dd") if self.has_date.isChecked() else "",
                    subs=self.subs, labels=list(self.sel_labels))


class LabelDialog(QDialog):
    def __init__(self, parent, store):
        super().__init__(parent); self.store = store; self.setWindowTitle("Labels"); self.setMinimumWidth(420)
        self.v = QVBoxLayout(self); self.list_box = QVBoxLayout(); self.v.addLayout(self.list_box)
        r = QHBoxLayout(); self.inp = QLineEdit(); self.inp.setPlaceholderText("New label name"); self.inp.returnPressed.connect(self.add)
        self.color = LABEL_PALETTE[len(store.labels) % len(LABEL_PALETTE)]
        self.cbtn = QPushButton("Colour"); self.cbtn.clicked.connect(self.pick); ab = QPushButton("Add"); ab.setObjectName("primary"); ab.clicked.connect(self.add)
        r.addWidget(self.inp); r.addWidget(self.cbtn); r.addWidget(ab); self.v.addLayout(r)
        cb = QPushButton("Close"); cb.clicked.connect(self.accept); self.v.addWidget(cb)
        self.rebuild()

    def pick(self):
        c = QColorDialog.getColor(QColor(self.color), self, "Label colour")
        if c.isValid():
            self.color = c.name(); self.cbtn.setStyleSheet(f"background:{self.color};color:#fff;")

    def rebuild(self):
        while self.list_box.count():
            w = self.list_box.takeAt(0).widget()
            if w: w.deleteLater()
        if not self.store.labels:
            self.list_box.addWidget(QLabel("No labels yet — add one below, e.g. Client, Urgent, Personal.", objectName="muted"))
        for l in self.store.labels:
            row = QHBoxLayout(); w = QWidget(); w.setLayout(row)
            sw = QPushButton(); sw.setFixedSize(26, 22); sw.setStyleSheet(f"background:{l['color']};border-radius:6px;")
            sw.clicked.connect(lambda _, lab=l: self.recolor(lab))
            name = QLineEdit(l["name"]); name.editingFinished.connect(lambda lab=l, ed=name: self.rename(lab, ed.text()))
            n = self.usage(l["id"]); cnt = QLabel(f"{n} tasks", objectName="muted")
            dl = QPushButton("✕"); dl.setObjectName("ghost"); dl.clicked.connect(lambda _, lab=l: self.delete(lab))
            for x in (sw, name, cnt, dl):
                row.addWidget(x)
            self.list_box.addWidget(w)

    def usage(self, lid):
        n = 0
        for Y in self.store.db["years"].values():
            for kd in Y["kanban"].values():
                for arr in kd.values():
                    n += sum(1 for c in arr if lid in c.get("labels", []))
        return n

    def add(self):
        t = self.inp.text().strip()
        if not t: return
        self.store.labels.append({"id": "l_" + uid(), "name": t, "color": self.color}); self.inp.clear()
        self.color = LABEL_PALETTE[len(self.store.labels) % len(LABEL_PALETTE)]; self.cbtn.setStyleSheet("")
        self.store.save(); self.rebuild()

    def rename(self, lab, t):
        if t.strip() and t.strip() != lab["name"]:
            lab["name"] = t.strip(); self.store.save()

    def recolor(self, lab):
        c = QColorDialog.getColor(QColor(lab["color"]), self, "Label colour")
        if c.isValid():
            lab["color"] = c.name(); self.store.save(); self.rebuild()

    def delete(self, lab):
        if QMessageBox.question(self, "Delete label", f"Delete label \"{lab['name']}\"? It will be removed from {self.usage(lab['id'])} task(s).") != QMessageBox.Yes:
            return
        for Y in self.store.db["years"].values():
            for kd in Y["kanban"].values():
                for arr in kd.values():
                    for c in arr:
                        if "labels" in c: c["labels"] = [x for x in c["labels"] if x != lab["id"]]
        self.store.labels.remove(lab); self.store.save(); self.rebuild()


class ColumnDialog(QDialog):
    def __init__(self, parent, store, col_id=None):
        super().__init__(parent); self.store = store; self.col = store.col(col_id) if col_id else None; self.result_action = None
        self.setWindowTitle("Column settings" if self.col else "New column"); self.setMinimumWidth(400)
        v = QVBoxLayout(self)
        self.name = QLineEdit(self.col["label"] if self.col else ""); self.name.setPlaceholderText("e.g. Review, Blocked, Waiting")
        self.icon = QLineEdit(self.col.get("icon", "") if self.col else "📌"); self.icon.setMaxLength(4); self.icon.setFixedWidth(60)
        self.color = self.col["color"] if self.col else LABEL_PALETTE[len(store.columns) % len(LABEL_PALETTE)]
        self.cbtn = QPushButton("Colour"); self.cbtn.setStyleSheet(f"background:{self.color};color:#fff;"); self.cbtn.clicked.connect(self.pick)
        r = QHBoxLayout(); r.addWidget(self.name); r.addWidget(self.icon); r.addWidget(self.cbtn); v.addWidget(QLabel("Name *")); v.addLayout(r)
        if self.col and self.col["id"] in BUILTIN:
            v.addWidget(QLabel("Built-in column: rename, recolour and move it, but it can't be deleted — the dashboard and month rollover depend on it.", objectName="muted", wordWrap=True))
        br = QHBoxLayout()
        if self.col:
            for txt, act in (("◀ Move left", "left"), ("Move right ▶", "right")):
                b = QPushButton(txt); b.clicked.connect(lambda _, a=act: self.finish(a)); br.addWidget(b)
            if self.col["id"] not in BUILTIN:
                d = QPushButton("Delete"); d.setObjectName("danger"); d.clicked.connect(lambda: self.finish("delete")); br.addWidget(d)
        br.addStretch(); sv = QPushButton("Save"); sv.setObjectName("primary"); sv.clicked.connect(lambda: self.finish("save")); cn = QPushButton("Cancel"); cn.clicked.connect(self.reject)
        br.addWidget(sv); br.addWidget(cn); v.addLayout(br)

    def pick(self):
        c = QColorDialog.getColor(QColor(self.color), self, "Column colour")
        if c.isValid():
            self.color = c.name(); self.cbtn.setStyleSheet(f"background:{self.color};color:#fff;")

    def finish(self, action):
        if action == "save" and not self.name.text().strip():
            self.name.setFocus(); return
        self.result_action = action; self.accept()


class ChecklistTaskDialog(QDialog):
    def __init__(self, parent, task):
        super().__init__(parent); self.setWindowTitle("Edit task"); self.setMinimumWidth(420)
        v = QVBoxLayout(self)
        self.title = QLineEdit(task["title"]); v.addWidget(QLabel("Title *")); v.addWidget(self.title)
        self.pri = QComboBox()
        for k, lab in PRI_LABEL.items(): self.pri.addItem(lab, k)
        self.pri.setCurrentIndex(self.pri.findData(task.get("priority", "med")))
        self.has_date = QCheckBox("Due date"); self.date = QDateEdit(QDate.currentDate()); self.date.setCalendarPopup(True); self.date.setDisplayFormat("yyyy-MM-dd")
        if task.get("dueDate"):
            self.has_date.setChecked(True); self.date.setDate(QDate.fromString(task["dueDate"], "yyyy-MM-dd"))
        self.date.setEnabled(self.has_date.isChecked()); self.has_date.toggled.connect(self.date.setEnabled)
        r = QHBoxLayout(); r.addWidget(QLabel("Priority")); r.addWidget(self.pri); r.addWidget(self.has_date); r.addWidget(self.date); v.addLayout(r)
        self.subs = [dict(s) for s in task.get("subs", [])]; self.sub_list = QListWidget(); self.sub_list.setFixedHeight(100); self.rebuild(); v.addWidget(QLabel("Sub-tasks")); v.addWidget(self.sub_list)
        sr = QHBoxLayout(); self.inp = QLineEdit(); self.inp.setPlaceholderText("Add sub-task… (Enter)"); self.inp.returnPressed.connect(self.add)
        rb = QPushButton("Remove selected"); rb.clicked.connect(self.remove); sr.addWidget(self.inp); sr.addWidget(rb); v.addLayout(sr)
        bb = QDialogButtonBox(QDialogButtonBox.Save | QDialogButtonBox.Cancel); bb.button(QDialogButtonBox.Save).setObjectName("primary")
        bb.accepted.connect(self.accept); bb.rejected.connect(self.reject); v.addWidget(bb)
        QShortcut(QKeySequence("Ctrl+Return"), self, activated=self.accept)

    def rebuild(self):
        self.sub_list.clear()
        for s in self.subs: self.sub_list.addItem(("✓ " if s.get("done") else "○ ") + s["text"])

    def add(self):
        t = self.inp.text().strip()
        if t: self.subs.append({"id": uid(), "text": t, "done": False}); self.inp.clear(); self.rebuild()

    def remove(self):
        r = self.sub_list.currentRow()
        if r >= 0: del self.subs[r]; self.rebuild()

    def accept(self):
        if not self.title.text().strip(): self.title.setFocus(); return
        super().accept()

    def apply(self, task):
        task.update(title=self.title.text().strip(), priority=self.pri.currentData(),
                    dueDate=self.date.date().toString("yyyy-MM-dd") if self.has_date.isChecked() else "", subs=self.subs, updatedAt=now_iso())


class ExportDialog(QDialog):
    def __init__(self, parent, store):
        super().__init__(parent); self.store = store; self.setWindowTitle("Export to Excel"); self.setMinimumWidth(460)
        v = QVBoxLayout(self); v.addWidget(QLabel(f"Select months to export ({store.year})", objectName="sectitle"))
        g = QGridLayout(); self.checks = []
        for i, m in enumerate(MONTHS):
            cb = QCheckBox(m); cb.setChecked(i == store.month)
            kd = store.kanban(i); has = any(kd.get(c["id"]) for c in store.columns) or any(gr["tasks"] for gr in store.checklist(i))
            if has: cb.setText(m + " •")
            g.addWidget(cb, i // 4, i % 4); self.checks.append(cb)
        v.addLayout(g)
        r = QHBoxLayout()
        for txt, fn in (("All", lambda: [c.setChecked(True) for c in self.checks]), ("None", lambda: [c.setChecked(False) for c in self.checks]),
                        ("With data", lambda: [c.setChecked("•" in c.text()) for c in self.checks])):
            b = QPushButton(txt); b.clicked.connect(fn); r.addWidget(b)
        r.addStretch(); v.addLayout(r)
        self.inc_board = QCheckBox("Board tasks (with sub-tasks)"); self.inc_board.setChecked(True)
        self.inc_cl = QCheckBox("Checklist groups & tasks"); self.inc_cl.setChecked(True)
        self.inc_sum = QCheckBox("Month-wise summary sheet"); self.inc_sum.setChecked(True)
        self.inc_done = QCheckBox("Include completed tasks"); self.inc_done.setChecked(True)
        for cb in (self.inc_board, self.inc_cl, self.inc_sum, self.inc_done): v.addWidget(cb)
        bb = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel); bb.button(QDialogButtonBox.Ok).setText("Export…"); bb.button(QDialogButtonBox.Ok).setObjectName("primary")
        bb.accepted.connect(self.accept); bb.rejected.connect(self.reject); v.addWidget(bb)

    def months(self):
        return [i for i, c in enumerate(self.checks) if c.isChecked()]


def export_excel(store, path, months, inc_board, inc_cl, inc_sum, inc_done):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook(); wb.remove(wb.active)
    hdr = Font(bold=True, color="FFFFFF"); fill = PatternFill("solid", fgColor="2B3563"); title = Font(bold=True, size=13)
    lname = lambda c: ", ".join(l["name"] for l in store.labels if l["id"] in c.get("labels", []))
    if inc_board:
        for mi in months:
            ws = wb.create_sheet(f"{MONTHS[mi]}-Board"); ws.append([f"{APP_NAME} — Board Tasks — {MONTHS[mi]} {store.year}"]); ws["A1"].font = title
            kd = store.kanban(mi)
            for c in store.columns:
                cards = [x for x in kd.get(c["id"], []) if inc_done or not x.get("done")]
                if not cards or (c["id"] == "done" and not inc_done): continue
                ws.append([]); ws.append([f"{c.get('icon', '')} {c['label'].upper()}"]); ws.cell(ws.max_row, 1).font = Font(bold=True, color=c["color"].lstrip("#"))
                ws.append(["#", "Task Title", "Description", "Labels", "Priority", "Due Date", "Status", "Migrated From", "Sub-tasks", "Done/Total", "Created"])
                for cell in ws[ws.max_row]: cell.font = hdr; cell.fill = fill
                for n, x in enumerate(cards, 1):
                    subs = x.get("subs", []); sd = sum(1 for s in subs if s.get("done"))
                    ws.append([n, x["title"], x.get("desc", ""), lname(x), PRI_LABEL[x.get("priority", "med")], x.get("dueDate", ""),
                               "Completed" if x.get("done") or c["id"] == "done" else c["label"], x.get("migratedFrom", ""),
                               " | ".join(("[✓] " if s.get("done") else "[ ] ") + s["text"] for s in subs), f"{sd}/{len(subs)}" if subs else "", (x.get("createdAt") or "")[:10]])
            for col, w in zip("ABCDEFGHIJK", (4, 30, 28, 16, 10, 12, 12, 13, 36, 10, 12)): ws.column_dimensions[col].width = w
    if inc_cl:
        for mi in months:
            groups = store.checklist(mi)
            if not groups: continue
            ws = wb.create_sheet(f"{MONTHS[mi]}-Checklist"); ws.append([f"{APP_NAME} — Checklists — {MONTHS[mi]} {store.year}"]); ws["A1"].font = title
            for g in groups:
                tasks = [t for t in g["tasks"] if inc_done or not t.get("done")]
                if not tasks: continue
                done = sum(1 for t in g["tasks"] if t.get("done")); tot = len(g["tasks"])
                ws.append([]); ws.append([f"📂 {g['name'].upper()}", "", f"{done}/{tot} done ({round(done / tot * 100) if tot else 0}%)"]); ws.cell(ws.max_row, 1).font = Font(bold=True)
                ws.append(["#", "Task Title", "Priority", "Due Date", "Status", "Sub-tasks", "Done/Total"])
                for cell in ws[ws.max_row]: cell.font = hdr; cell.fill = fill
                for n, t in enumerate(tasks, 1):
                    subs = t.get("subs", []); sd = sum(1 for s in subs if s.get("done"))
                    ws.append([n, t["title"], PRI_LABEL[t.get("priority", "med")], t.get("dueDate", ""), "Done" if t.get("done") else "Pending",
                               " | ".join(("[✓] " if s.get("done") else "[ ] ") + s["text"] for s in subs), f"{sd}/{len(subs)}" if subs else ""])
            for col, w in zip("ABCDEFG", (4, 34, 10, 12, 11, 38, 10)): ws.column_dimensions[col].width = w
    if inc_sum:
        ws = wb.create_sheet("Summary"); ws.append([f"{APP_NAME} — Month-wise Summary — {store.year}"]); ws["A1"].font = title; ws.append([])
        ws.append(["Month", "Year", "To Do", "Today's", "In Progress", "Completed", "Total", "Completion %", "CL Groups", "CL Tasks", "CL Done", "CL %"])
        for cell in ws[ws.max_row]: cell.font = hdr; cell.fill = fill
        ip = [c["id"] for c in store.columns if c["id"] not in ("todo", "today", "done")]
        for mi in range(12):
            kd = store.kanban(mi); todo, today, done = len(kd.get("todo", [])), len(kd.get("today", [])), len(kd.get("done", []))
            wip = sum(len(kd.get(i, [])) for i in ip); tot = todo + today + wip + done
            gs = store.checklist(mi); ct = sum(len(g["tasks"]) for g in gs); cd = sum(1 for g in gs for t in g["tasks"] if t.get("done"))
            ws.append([MONTHS[mi], store.year, todo, today, wip, done, tot, f"{round(done / tot * 100) if tot else 0}%", len(gs), ct, cd, f"{round(cd / ct * 100) if ct else 0}%"])
    if not wb.sheetnames: wb.create_sheet("Empty").append(["Nothing selected"])
    wb.save(path)


# ═══════════════════════════════════════════════════════════════════════════
#  MAIN WINDOW
# ═══════════════════════════════════════════════════════════════════════════
class MainWindow(QMainWindow):
    def __init__(self, store):
        super().__init__(); self.store = store; self.setWindowTitle(APP_NAME); self.resize(1400, 880); self.setMinimumSize(980, 620)
        self.detail_ref = None; self.filter_q = ""; self.filter_pri = ""; self.filter_labels = set(); self.cl_filter = "all"; self.cl_q = ""
        self.settings = read_settings(); self.appearance = {**dict(theme="dark", font="", size=13, text="", accent="", bg="", surface=""), **self.settings.get("appearance", {}), **self.store.db["meta"].get("appearance_py", {})}
        self._build(); self.apply_appearance(); self._tray(); self.show_page("dashboard")
        moved = self.store.auto_migrate()
        if moved:
            self.toast(f"🔄 {moved} incomplete task(s) auto-moved to the next month"); self.refresh()
        self.backup_timer = QTimer(self); self.backup_timer.timeout.connect(lambda: self.store.backup_now()); self.backup_timer.start(30 * 60 * 1000)
        QShortcut(QKeySequence("Ctrl+N"), self, activated=lambda: self.add_task())
        QShortcut(QKeySequence("Ctrl+Shift+B"), self, activated=self.backup_now_ui)

    # ---------- layout ----------
    def _build(self):
        root = QWidget(); self.setCentralWidget(root); h = QHBoxLayout(root); h.setContentsMargins(0, 0, 0, 0); h.setSpacing(0)
        sb = QFrame(); sb.setObjectName("sidebar"); sb.setFixedWidth(220); sv = QVBoxLayout(sb); sv.setContentsMargins(0, 20, 0, 16); sv.setSpacing(2)
        b = QLabel(APP_NAME); b.setObjectName("brand"); b.setContentsMargins(20, 0, 0, 0); sv.addWidget(b)
        bs = QLabel("PLAN · DO · TRACK"); bs.setObjectName("brandSub"); bs.setContentsMargins(20, 0, 0, 14); sv.addWidget(bs)
        sv.addWidget(QLabel("WORKSPACE", objectName="navSection", indent=20)); self.nav_btns = {}
        for key, txt in (("dashboard", "◈  Dashboard"), ("kanban", "⊞  Task Board"), ("checklist", "☑  Task Checklist")):
            nb = QPushButton(txt); nb.setObjectName("nav"); nb.setCheckable(True); nb.clicked.connect(lambda _, k=key: self.show_page(k)); sv.addWidget(nb); self.nav_btns[key] = nb
        sv.addSpacing(10); sv.addWidget(QLabel("TOOLS", objectName="navSection", indent=20))
        self.roll_btn = QPushButton("🔄  Month Rollover"); self.roll_btn.setObjectName("nav"); self.roll_btn.clicked.connect(self.open_rollover); sv.addWidget(self.roll_btn)
        for txt, fn in (("💾  Backup && Restore", self.open_backup), ("🎨  Appearance", self.open_appearance)):
            nb = QPushButton(txt); nb.setObjectName("nav"); nb.clicked.connect(fn); sv.addWidget(nb)
        sv.addStretch()
        yr = QHBoxLayout(); yr.setContentsMargins(16, 0, 16, 6)
        pb = QPushButton("◀"); pb.setObjectName("ghost"); pb.clicked.connect(lambda: self.switch_year(-1)); self.year_lbl = QLabel(str(self.store.year)); self.year_lbl.setAlignment(Qt.AlignCenter); self.year_lbl.setStyleSheet("font-weight:700;font-size:15px;")
        nb2 = QPushButton("▶"); nb2.setObjectName("ghost"); nb2.clicked.connect(lambda: self.switch_year(1)); yr.addWidget(pb); yr.addWidget(self.year_lbl, 1); yr.addWidget(nb2); sv.addLayout(yr)
        mg = QGridLayout(); mg.setContentsMargins(16, 0, 16, 0); mg.setSpacing(4); self.month_btns = []
        for i, m in enumerate(MONTHS):
            mb = QPushButton(m); mb.setObjectName("month"); mb.setCheckable(True); mb.clicked.connect(lambda _, i=i: self.switch_month(i)); mg.addWidget(mb, i // 3, i % 3); self.month_btns.append(mb)
        sv.addLayout(mg); h.addWidget(sb)
        main = QVBoxLayout(); main.setContentsMargins(0, 0, 0, 0); main.setSpacing(0)
        tb = QFrame(); tb.setObjectName("topbar"); th = QHBoxLayout(tb); th.setContentsMargins(24, 12, 24, 12)
        self.page_title = QLabel("Dashboard"); self.page_title.setObjectName("pageTitle"); self.page_month = QLabel(); self.page_month.setObjectName("pageMonth"); self.page_month.setMargin(6)
        th.addWidget(self.page_title); th.addWidget(self.page_month); th.addStretch()
        self.theme_btn = QPushButton("🌙"); self.theme_btn.setToolTip("Toggle dark/light"); self.theme_btn.clicked.connect(self.toggle_theme)
        ex = QPushButton("⬇ Export Excel"); ex.clicked.connect(self.export_excel); ad = QPushButton("+ Add Task"); ad.setObjectName("primary"); ad.clicked.connect(lambda: self.add_task())
        th.addWidget(self.theme_btn); th.addWidget(ex); th.addWidget(ad); main.addWidget(tb)
        self.stack = QStackedWidget(); main.addWidget(self.stack, 1)
        self.pages = {}
        for key, build in (("dashboard", self._build_dashboard), ("kanban", self._build_board), ("checklist", self._build_checklist)):
            self.pages[key] = build(); self.stack.addWidget(self.pages[key])
        w = QWidget(); w.setLayout(main); h.addWidget(w, 1)
        # detail dock
        self.dock = QDockWidget("Task", self); self.dock.setAllowedAreas(Qt.RightDockWidgetArea); self.dock.setFeatures(QDockWidget.DockWidgetClosable); self.dock.setMinimumWidth(440)
        self.detail = QFrame(); self.detail.setObjectName("detail"); self.detail_layout = QVBoxLayout(self.detail); self.dock.setWidget(self.detail); self.addDockWidget(Qt.RightDockWidgetArea, self.dock); self.dock.hide()
        self.dock.visibilityChanged.connect(lambda vis: setattr(self, "detail_ref", None) if not vis else None)
        self.toast_lbl = QLabel(self); self.toast_lbl.setStyleSheet("background:#13161f;color:#eef0f8;border:1px solid #3ddbbf;border-radius:10px;padding:8px 16px;"); self.toast_lbl.hide()

    def _scroll(self, inner):
        sa = QScrollArea(); sa.setWidgetResizable(True); sa.setWidget(inner); return sa

    def _build_dashboard(self):
        w = QWidget(); v = QVBoxLayout(w); v.setContentsMargins(24, 20, 24, 20); v.setSpacing(16)
        self.kpi_row = QHBoxLayout(); v.addLayout(self.kpi_row)
        row = QHBoxLayout(); v.addLayout(row, 1)
        p1 = QFrame(); p1.setObjectName("panel"); l1 = QVBoxLayout(p1); l1.addWidget(QLabel("Recent Activity", objectName="sectitle")); self.activity_box = QVBoxLayout(); l1.addLayout(self.activity_box); l1.addStretch(); row.addWidget(p1, 3)
        p2 = QFrame(); p2.setObjectName("panel"); l2 = QVBoxLayout(p2); l2.addWidget(QLabel("Tasks / Month", objectName="sectitle"))
        self.chart = BarChart(lambda: [sum(len(self.store.kanban(i).get(c["id"], [])) for c in self.store.columns) for i in range(12)], "#6c8aff"); l2.addWidget(self.chart); l2.addStretch(); row.addWidget(p2, 2)
        return self._scroll(w)

    def _build_board(self):
        w = QWidget(); v = QVBoxLayout(w); v.setContentsMargins(24, 14, 24, 14); v.setSpacing(12)
        tbar = QHBoxLayout(); self.search = QLineEdit(); self.search.setPlaceholderText("Search tasks…"); self.search.setMinimumWidth(160); self.search.setMaximumWidth(260); self.search.textChanged.connect(self.on_search)
        self.pri_filter = QComboBox(); self.pri_filter.setMinimumWidth(140); self.pri_filter.addItem("All priorities", "")
        for k, lab in PRI_LABEL.items(): self.pri_filter.addItem(lab, k)
        self.pri_filter.currentIndexChanged.connect(lambda: (setattr(self, "filter_pri", self.pri_filter.currentData()), self.render_board()))
        self.label_bar = QHBoxLayout(); tbar.addWidget(self.search); tbar.addWidget(self.pri_filter); tbar.addLayout(self.label_bar); tbar.addStretch()
        lb = QPushButton("🏷 Labels"); lb.clicked.connect(self.manage_labels); cb = QPushButton("+ Column"); cb.clicked.connect(lambda: self.column_dialog(None)); tbar.addWidget(lb); tbar.addWidget(cb); v.addLayout(tbar)
        self.board_host = QWidget(); self.board_layout = QHBoxLayout(self.board_host); self.board_layout.setContentsMargins(0, 0, 0, 0); self.board_layout.setSpacing(14); self.board_layout.setAlignment(Qt.AlignLeft | Qt.AlignTop)
        sa = QScrollArea(); sa.setWidgetResizable(True); sa.setWidget(self.board_host); v.addWidget(sa, 1)
        return w

    def _build_checklist(self):
        w = QWidget(); v = QVBoxLayout(w); v.setContentsMargins(24, 14, 24, 14); v.setSpacing(12)
        tb = QHBoxLayout(); self.cl_btns = {}
        for k, t in (("all", "All"), ("pending", "Pending"), ("done", "Done")):
            b = QPushButton(t); b.setCheckable(True); b.setChecked(k == "all"); b.clicked.connect(lambda _, k=k: self.set_cl_filter(k)); tb.addWidget(b); self.cl_btns[k] = b
        self.cl_search = QLineEdit(); self.cl_search.setPlaceholderText("Search tasks…"); self.cl_search.setFixedWidth(240); self.cl_search.textChanged.connect(lambda t: (setattr(self, "cl_q", t), self.render_checklist()))
        tb.addWidget(self.cl_search); tb.addStretch(); ng = QPushButton("+ New Group"); ng.setObjectName("primary"); ng.clicked.connect(self.add_group); tb.addWidget(ng); v.addLayout(tb)
        self.cl_host = QWidget(); self.cl_layout = QVBoxLayout(self.cl_host); self.cl_layout.setAlignment(Qt.AlignTop); self.cl_layout.setSpacing(12)
        v.addWidget(self._scroll(self.cl_host), 1); return w

    # ---------- tray ----------
    def _tray(self):
        pm = QPixmap(64, 64); pm.fill(Qt.transparent); p = QPainter(pm); p.setRenderHint(QPainter.Antialiasing); p.setBrush(QColor("#6c8aff")); p.setPen(Qt.NoPen); p.drawRoundedRect(4, 4, 56, 56, 14, 14)
        p.setPen(QColor("#ffffff")); f = QFont(); f.setBold(True); f.setPixelSize(34); p.setFont(f); p.drawText(pm.rect(), Qt.AlignCenter, "T"); p.end()
        self.setWindowIcon(QIcon(pm)); self.tray = QSystemTrayIcon(QIcon(pm), self); m = QMenu()
        m.addAction(f"Open {APP_NAME}", self.show_window); m.addSeparator(); m.addAction("Backup Data Now", self.backup_now_ui)
        m.addAction("Open Backup Folder", lambda: self.open_path(get_backup_dir())); m.addAction("Open Data Folder", lambda: self.open_path(USER_DATA)); m.addSeparator(); m.addAction("Quit", self.quit_app)
        self.tray.setContextMenu(m); self.tray.activated.connect(lambda r: self.show_window() if r == QSystemTrayIcon.DoubleClick else None); self.tray.setToolTip(APP_NAME); self.tray.show()
        self.quitting = False; self._last_close = 0; self._hint_shown = False

    def show_window(self):
        self.show(); self.raise_(); self.activateWindow()

    def quit_app(self):
        self.quitting = True; self.store.save(); self.store.backup_now(); QApplication.quit()

    def closeEvent(self, e):
        if self.quitting: e.accept(); return
        if time.time() - self._last_close < 3: self.quitting = True; e.accept(); QApplication.quit(); return
        self._last_close = time.time(); e.ignore(); self.hide()
        if not self._hint_shown:
            self._hint_shown = True; self.tray.showMessage(APP_NAME, "Still running in the system tray. Right-click the tray icon and choose Quit to exit.")

    @staticmethod
    def open_path(p):
        try:
            if sys.platform.startswith("win"): os.startfile(p)
            elif sys.platform == "darwin": os.system(f'open "{p}"')
            else: os.system(f'xdg-open "{p}" &')
        except Exception:
            pass

    # ---------- helpers ----------
    def toast(self, msg, ms=2600):
        self.toast_lbl.setText(msg); self.toast_lbl.adjustSize(); self._place_toast(); self.toast_lbl.show(); self.toast_lbl.raise_()
        QTimer.singleShot(ms, self.toast_lbl.hide)

    def resizeEvent(self, e):
        super().resizeEvent(e)
        if self.toast_lbl.isVisible(): self._place_toast()

    def _place_toast(self):
        c = self.centralWidget().geometry()
        self.toast_lbl.move(c.right() - self.toast_lbl.width() - 24, c.bottom() - self.toast_lbl.height() - 24)

    def show_page(self, key):
        self.page = key
        for k, b in self.nav_btns.items(): b.setChecked(k == key)
        self.stack.setCurrentWidget(self.pages[key]); self.page_title.setText({"dashboard": "Dashboard", "kanban": "Task Board", "checklist": "Task Checklist"}[key]); self.refresh()

    def refresh(self):
        s = self.store; self.page_month.setText(f"{MONTHS_LONG[s.month]} {s.year}"); self.year_lbl.setText(str(s.year))
        for i, b in enumerate(self.month_btns): b.setChecked(i == s.month)
        kd = s.kanban(); n = sum(len(kd.get(c["id"], [])) for c in s.columns); ncl = sum(len(g["tasks"]) for g in s.checklist())
        self.nav_btns["kanban"].setText(f"⊞  Task Board   ({n})"); self.nav_btns["checklist"].setText(f"☑  Task Checklist   ({ncl})")
        self.roll_btn.setText("🔄  Month Rollover" + ("  ⚠" if s.pending_rollover() else ""))
        {"dashboard": self.render_dashboard, "kanban": self.render_board, "checklist": self.render_checklist}[self.page]()
        if self.detail_ref: self.render_detail()

    def switch_month(self, i):
        self.store.month = i; self.close_detail(); self.refresh()

    def switch_year(self, d):
        ys = self.store.years(); ry = dt.date.today().year; lo, hi = min(ys[0] if ys else ry, ry - 1), max(ys[-1] if ys else ry, ry + 1)
        ny = self.store.year + d
        if lo <= ny <= hi: self.store.year = ny; self.store.ensure_year(ny); self.close_detail(); self.refresh()
        else: self.toast(f"Years available: {lo} – {hi}")

    @staticmethod
    def _clear(layout):
        while layout.count():
            it = layout.takeAt(0)
            if it.widget(): it.widget().deleteLater()
            elif it.layout(): MainWindow._clear(it.layout())

    # ---------- dashboard ----------
    def render_dashboard(self):
        s = self.store; kd = s.kanban(); self._clear(self.kpi_row)
        todo, today, done = len(kd.get("todo", [])), len(kd.get("today", [])), len(kd.get("done", []))
        wip = sum(len(kd.get(c["id"], [])) for c in s.columns if c["id"] not in ("todo", "today", "done")); total = todo + today + wip + done
        pct = round(done / total * 100) if total else 0
        overdue = sum(1 for c in s.columns if c["id"] != "done" for x in kd.get(c["id"], []) if is_overdue(x.get("dueDate"), x.get("done")))
        for lab, val, sub, color in (("TOTAL TASKS", total, f"{overdue} overdue" if overdue else "This month", "#6c8aff"), ("TODAY'S FOCUS", today, "Tasks for today", "#ffb347"),
                                     ("IN PROGRESS", wip, "Currently working", "#b48aff"), ("COMPLETED", done, f"{pct}% of month", "#3ddbbf")):
            f = QFrame(); f.setObjectName("kpi"); fv = QVBoxLayout(f); fv.addWidget(QLabel(lab, objectName="kpiLabel")); fv.addWidget(QLabel(str(val), objectName="kpiVal"))
            sl = QLabel(sub); sl.setObjectName("muted"); sl.setStyleSheet(f"color:{'#ff6584' if 'overdue' in sub else '#5c6382'}"); fv.addWidget(sl); f.setStyleSheet(f"QFrame#kpi{{border-bottom:3px solid {color};}}"); self.kpi_row.addWidget(f)
        self._clear(self.activity_box)
        cards = [(x, c) for c in s.columns for x in kd.get(c["id"], [])]
        cards.sort(key=lambda t: t[0].get("updatedAt") or t[0].get("createdAt") or "", reverse=True)
        if not cards: self.activity_box.addWidget(QLabel("No tasks this month yet.", objectName="muted"))
        for x, c in cards[:6]:
            b = QPushButton(("✓ " if x.get("done") else "○ ") + x["title"] + f"   ·  {c['label']} · {PRI_LABEL[x.get('priority', 'med')]}"); b.setObjectName("addk")
            b.clicked.connect(lambda _, cid=x["id"], col=c["id"]: (self.show_page("kanban"), self.open_detail(cid, col))); self.activity_box.addWidget(b)
        self.chart.update()

    # ---------- board ----------
    def on_search(self, t):
        self.filter_q = t.lower(); self.render_board()

    def card_matches(self, c):
        if self.filter_pri and c.get("priority") != self.filter_pri: return False
        if self.filter_labels and not self.filter_labels <= set(c.get("labels", [])): return False
        if self.filter_q:
            hay = " ".join([c.get("title", ""), c.get("desc", "")] + [s["text"] for s in c.get("subs", [])]).lower()
            if self.filter_q not in hay: return False
        return True

    def render_board(self):
        s = self.store; self._clear(self.label_bar)
        for l in s.labels:
            b = QPushButton(l["name"]); b.setCheckable(True); b.setChecked(l["id"] in self.filter_labels)
            b.setStyleSheet(f"QPushButton{{border-color:{l['color']};color:{l['color']};padding:3px 10px;}}QPushButton:checked{{background:{l['color']};color:#fff;}}")
            b.setSizePolicy(QSizePolicy.Fixed, QSizePolicy.Fixed)
            b.toggled.connect(lambda on, lid=l["id"]: (self.filter_labels.add(lid) if on else self.filter_labels.discard(lid), self.render_board())); self.label_bar.addWidget(b)
        self._clear(self.board_layout); kd = s.kanban()
        for col in s.columns:
            cw = QWidget(); cw.setFixedWidth(280); cv = QVBoxLayout(cw); cv.setContentsMargins(0, 0, 0, 0); cv.setSpacing(8)
            head = QFrame(); head.setObjectName("colhead"); hh = QHBoxLayout(head); hh.setContentsMargins(12, 8, 8, 8)
            dot = QLabel("●"); dot.setStyleSheet(f"color:{col['color']};"); t = QLabel(f"{col.get('icon', '')} {col['label']}"); t.setStyleSheet("font-weight:700;")
            pal = self.palette_dict(); cnt = badge(str(len(kd.get(col["id"], []))), pal["text2"], pal["s4"]); mb = QPushButton("…"); mb.setObjectName("ghost"); mb.setFixedWidth(26); mb.clicked.connect(lambda _, cid=col["id"]: self.column_dialog(cid))
            hh.addWidget(dot); hh.addWidget(t); hh.addStretch(); hh.addWidget(cnt); hh.addWidget(mb); cv.addWidget(head)
            lst = CardList(col["id"]); lst.dropped.connect(self.on_drop)
            cards = [c for c in kd.get(col["id"], []) if self.card_matches(c)]
            if not kd.get(col["id"]):
                it = QListWidgetItem("📭  Empty"); it.setFlags(Qt.NoItemFlags); it.setTextAlignment(Qt.AlignCenter); lst.addItem(it)
            elif not cards:
                it = QListWidgetItem("🔍  No matches"); it.setFlags(Qt.NoItemFlags); it.setTextAlignment(Qt.AlignCenter); lst.addItem(it)
            for c in cards:
                it = QListWidgetItem(); it.setData(Qt.UserRole, c["id"]); w = CardWidget(s, c, col); it.setSizeHint(w.sizeHint() + QSize(0, 6)); lst.addItem(it); lst.setItemWidget(it, w)
                w.clicked.connect(self.open_detail); w.edit.connect(self.edit_task); w.delete.connect(self.delete_task); w.toggle.connect(self.toggle_done)
            lst.fit_items(); cv.addWidget(lst, 1); ab = QPushButton("+ Add Task"); ab.setObjectName("addk"); ab.clicked.connect(lambda _, cid=col["id"]: self.add_task(cid)); cv.addWidget(ab); self.board_layout.addWidget(cw)
        addc = QPushButton("+ Add column"); addc.setObjectName("addk"); addc.setFixedWidth(200); addc.clicked.connect(lambda: self.column_dialog(None)); self.board_layout.addWidget(addc, 0, Qt.AlignTop)

    def on_drop(self, cid, from_col, to_col, before):
        if before == cid: return
        card = self.store.move_card(cid, from_col, to_col, before)
        if card:
            self.store.save(); self.refresh()
            if self.detail_ref and self.detail_ref[0] == cid: self.detail_ref = (cid, to_col); self.render_detail()
            if to_col == "done" and from_col != "done": self.toast("✅ Task completed!")

    def add_task(self, col_id="todo"):
        d = TaskDialog(self, self.store, None, col_id)
        if d.exec() != QDialog.Accepted: return
        v = d.values(); card = dict(id=uid(), title=v["title"], desc=v["desc"], priority=v["priority"], dueDate=v["dueDate"], done=v["col"] == "done",
                                    subs=v["subs"], labels=v["labels"], comments=[], activity=[], createdAt=now_iso(), month=MONTHS[self.store.month], year=self.store.year)
        self.store.log(card, f"Created in {self.store.col(v['col'])['label']}"); self.store.kanban().setdefault(v["col"], []).insert(0, card); self.store.save(); self.refresh(); self.toast("Task added!")

    def edit_task(self, cid, col_id):
        card, col_id = self.store.find_card(cid, col_id)
        if not card: return
        d = TaskDialog(self, self.store, card, col_id)
        if d.exec() != QDialog.Accepted: return
        v = d.values(); changes = [n for n, k in (("title", "title"), ("description", "desc"), ("priority", "priority"), ("due date", "dueDate")) if (card.get(k) or "") != v[k]]
        card.update(title=v["title"], desc=v["desc"], priority=v["priority"], dueDate=v["dueDate"], subs=v["subs"], labels=v["labels"], updatedAt=now_iso())
        if changes: self.store.log(card, "Edited " + ", ".join(changes))
        if v["col"] != col_id: self.store.move_card(cid, col_id, v["col"])
        elif v["col"] == "done": card["done"] = True
        if self.detail_ref and self.detail_ref[0] == cid: self.detail_ref = (cid, v["col"])
        self.store.save(); self.refresh(); self.toast("Task updated")

    def delete_task(self, cid, col_id):
        if QMessageBox.question(self, "Delete task", "Delete this task? This cannot be undone.") != QMessageBox.Yes: return
        kd = self.store.kanban(); kd[col_id] = [c for c in kd.get(col_id, []) if c["id"] != cid]
        if self.detail_ref and self.detail_ref[0] == cid: self.close_detail()
        self.store.save(); self.refresh()

    def toggle_done(self, cid, col_id):
        to = "done" if col_id != "done" else "todo"
        if self.store.move_card(cid, col_id, to):
            self.store.save()
            if self.detail_ref and self.detail_ref[0] == cid: self.detail_ref = (cid, to)
            self.refresh(); self.toast("✅ Moved to Completed!" if to == "done" else "↩ Moved to To Do")

    def manage_labels(self):
        LabelDialog(self, self.store).exec(); self.filter_labels &= {l["id"] for l in self.store.labels}; self.refresh()

    def column_dialog(self, col_id):
        d = ColumnDialog(self, self.store, col_id)
        if d.exec() != QDialog.Accepted: return
        s = self.store; cols = s.columns; a = d.result_action
        if a == "save":
            if d.col: d.col.update(label=d.name.text().strip(), icon=d.icon.text().strip(), color=d.color); self.toast("Column updated")
            else: cols.append({"id": "c_" + uid(), "label": d.name.text().strip(), "icon": d.icon.text().strip(), "color": d.color}); [s.ensure_year(y) for y in s.years()]; self.toast("Column added")
        elif a in ("left", "right"):
            i = cols.index(d.col); j = i + (-1 if a == "left" else 1)
            if 0 <= j < len(cols): cols[i], cols[j] = cols[j], cols[i]
        elif a == "delete":
            n = sum(len(kd.get(d.col["id"], [])) for Y in s.db["years"].values() for kd in Y["kanban"].values())
            if QMessageBox.question(self, "Delete column", f"Delete column \"{d.col['label']}\"?" + (f" Its {n} task(s) across all months will move to To Do." if n else "")) != QMessageBox.Yes: return
            for Y in s.db["years"].values():
                for kd in Y["kanban"].values():
                    arr = kd.pop(d.col["id"], [])
                    for c in arr: c["done"] = False; s.log(c, f"Column \"{d.col['label']}\" deleted — moved to To Do")
                    kd["todo"] = arr + kd.get("todo", [])
            cols.remove(d.col)
        s.save(); self.refresh()

    # ---------- detail panel ----------
    def open_detail(self, cid, col_id):
        card, col_id = self.store.find_card(cid, col_id)
        if not card: return
        self.detail_ref = (cid, col_id); self.render_detail(); self.dock.show()

    def close_detail(self):
        self.detail_ref = None; self.dock.hide()

    def _dcard(self):
        if not self.detail_ref: return None, None
        return self.store.find_card(*self.detail_ref)

    def render_detail(self):
        card, col_id = self._dcard()
        if not card: self.close_detail(); return
        s = self.store; col = s.col(col_id); done = card.get("done") or col_id == "done"
        self._clear(self.detail_layout); L = self.detail_layout; L.setContentsMargins(0, 0, 0, 0)
        inner = QWidget(); v = QVBoxLayout(inner); v.setContentsMargins(20, 16, 20, 16); v.setSpacing(12)
        top = QHBoxLayout(); top.addWidget(badge(f"{col.get('icon', '')} {col['label']}", "#fff", col["color"])); top.addWidget(badge(PRI_LABEL[card.get("priority", "med")], PRI_COLOR[card.get("priority", "med")]))
        if card.get("dueDate"): od = is_overdue(card["dueDate"], done); top.addWidget(badge(("⚠ Overdue " if od else "📅 ") + fmt_date(card["dueDate"]), "#ff6584" if od else "#8890b0"))
        top.addStretch(); cl = QPushButton("✕"); cl.setObjectName("ghost"); cl.clicked.connect(self.close_detail); top.addWidget(cl); v.addLayout(top)
        title = QLineEdit(card["title"]); title.setStyleSheet("font-size:17px;font-weight:700;background:transparent;border-color:transparent;"); title.editingFinished.connect(lambda: self.dset("title", title.text())); v.addWidget(title)
        g = QGridLayout(); colc = QComboBox()
        for c in s.columns: colc.addItem(f"{c.get('icon', '')} {c['label']}", c["id"])
        colc.setCurrentIndex(colc.findData(col_id)); colc.currentIndexChanged.connect(lambda: self.dmove(colc.currentData()))
        pri = QComboBox()
        for k, lab in PRI_LABEL.items(): pri.addItem(lab, k)
        pri.setCurrentIndex(pri.findData(card.get("priority", "med"))); pri.currentIndexChanged.connect(lambda: self.dset("priority", pri.currentData()))
        dcb = QCheckBox("Due"); de = QDateEdit(QDate.currentDate()); de.setCalendarPopup(True); de.setDisplayFormat("yyyy-MM-dd")
        if card.get("dueDate"): dcb.setChecked(True); de.setDate(QDate.fromString(card["dueDate"], "yyyy-MM-dd"))
        de.setEnabled(dcb.isChecked()); dcb.toggled.connect(lambda on: self.dset("dueDate", de.date().toString("yyyy-MM-dd") if on else "")); de.dateChanged.connect(lambda d: self.dset("dueDate", d.toString("yyyy-MM-dd")) if dcb.isChecked() else None)
        for i, (lab, w) in enumerate((("Column", colc), ("Priority", pri))): g.addWidget(QLabel(lab, objectName="muted"), 0, i); g.addWidget(w, 1, i)
        dl = QHBoxLayout(); dl.addWidget(dcb); dl.addWidget(de); g.addWidget(QLabel("Due date", objectName="muted"), 0, 2); g.addLayout(dl, 1, 2); v.addLayout(g)
        v.addWidget(QLabel("LABELS", objectName="navSection")); lr = QHBoxLayout()
        for l in s.labels:
            b = QPushButton(l["name"]); b.setCheckable(True); b.setChecked(l["id"] in card.get("labels", [])); b.setStyleSheet(f"QPushButton{{border-color:{l['color']};color:{l['color']};padding:3px 10px;}}QPushButton:checked{{background:{l['color']};color:#fff;}}")
            b.toggled.connect(lambda on, lid=l["id"]: self.dlabel(lid, on)); lr.addWidget(b)
        if not s.labels: lr.addWidget(QLabel("No labels yet", objectName="muted"))
        lr.addStretch(); mg = QPushButton("Manage"); mg.clicked.connect(self.manage_labels); lr.addWidget(mg); v.addLayout(lr)
        v.addWidget(QLabel("DESCRIPTION", objectName="navSection")); desc = QTextEdit(card.get("desc", "")); desc.setPlaceholderText("Add details, links, notes…"); desc.setFixedHeight(80)
        desc.focusOutEvent = lambda e, te=desc: (QTextEdit.focusOutEvent(te, e), self.dset("desc", te.toPlainText())); v.addWidget(desc)
        subs = card.get("subs", []); sd = sum(1 for x in subs if x.get("done")); v.addWidget(QLabel(f"CHECKLIST  {sd}/{len(subs)}" if subs else "CHECKLIST", objectName="navSection"))
        for x in subs:
            r = QHBoxLayout(); cb = QCheckBox(x["text"]); cb.setChecked(bool(x.get("done"))); cb.toggled.connect(lambda on, sid=x["id"]: self.dsub(sid, on)); rm = QPushButton("✕"); rm.setObjectName("ghost"); rm.clicked.connect(lambda _, sid=x["id"]: self.dsub_remove(sid)); r.addWidget(cb, 1); r.addWidget(rm); v.addLayout(r)
        ar = QHBoxLayout(); si = QLineEdit(); si.setPlaceholderText("Add an item… (Enter)"); si.returnPressed.connect(lambda: self.dsub_add(si.text())); ab = QPushButton("Add"); ab.clicked.connect(lambda: self.dsub_add(si.text())); ar.addWidget(si); ar.addWidget(ab); v.addLayout(ar)
        comments = card.get("comments", []); v.addWidget(QLabel(f"COMMENTS  {len(comments) or ''}", objectName="navSection"))
        cr = QHBoxLayout(); ci = QTextEdit(); ci.setPlaceholderText("Write a comment… (Ctrl+Enter to post)"); ci.setFixedHeight(56); pb = QPushButton("Post"); pb.setObjectName("primary"); pb.clicked.connect(lambda: self.dcomment(ci.toPlainText()))
        QShortcut(QKeySequence("Ctrl+Return"), ci, activated=lambda: self.dcomment(ci.toPlainText())); cr.addWidget(ci, 1); cr.addWidget(pb, 0, Qt.AlignBottom); v.addLayout(cr)
        for m in reversed(comments):
            f = QFrame(); f.setObjectName("panel"); fl = QVBoxLayout(f); hr = QHBoxLayout(); hr.addWidget(QLabel(fmt_dt(m["at"]), objectName="muted")); hr.addStretch(); dbn = QPushButton("Delete"); dbn.setObjectName("ghost"); dbn.clicked.connect(lambda _, mid=m["id"]: self.dcomment_del(mid)); hr.addWidget(dbn); fl.addLayout(hr)
            t = QLabel(m["text"]); t.setWordWrap(True); fl.addWidget(t); v.addWidget(f)
        v.addWidget(QLabel("ACTIVITY", objectName="navSection"))
        for a in reversed(card.get("activity", [])[-20:]): v.addWidget(QLabel(f"{fmt_dt(a['at'])}   {a['text']}", objectName="muted2", wordWrap=True))
        v.addStretch(); sa = QScrollArea(); sa.setWidgetResizable(True); sa.setWidget(inner); L.addWidget(sa, 1)
        foot = QHBoxLayout(); foot.setContentsMargins(16, 8, 16, 12); mk = QPushButton("↩ Reopen" if done else "✓ Mark complete"); mk.setObjectName("" if done else "primary"); mk.clicked.connect(lambda: self.toggle_done(card["id"], col_id))
        ed = QPushButton("✎ Edit in form"); ed.clicked.connect(lambda: self.edit_task(card["id"], col_id)); dl2 = QPushButton("Delete"); dl2.setObjectName("danger"); dl2.clicked.connect(lambda: self.delete_task(card["id"], col_id))
        foot.addWidget(mk); foot.addWidget(ed); foot.addStretch(); foot.addWidget(dl2); L.addLayout(foot)

    def dset(self, field, val):
        card, col = self._dcard()
        if not card: return
        val = val.strip() if isinstance(val, str) else val
        if field == "title" and not val: self.render_detail(); return
        if (card.get(field) or "") == val: return
        card[field] = val; card["updatedAt"] = now_iso()
        self.store.log(card, {"title": "Title updated", "desc": "Description updated", "priority": f"Priority set to {PRI_LABEL.get(val, val)}", "dueDate": f"Due date set to {fmt_date(val)}" if val else "Due date removed"}[field])
        self.store.save(); self.refresh()

    def dmove(self, to):
        card, col = self._dcard()
        if not card or to == col: return
        self.store.move_card(card["id"], col, to); self.detail_ref = (card["id"], to); self.store.save(); self.refresh()

    def dlabel(self, lid, on):
        card, col = self._dcard()
        if not card: return
        ls = card.setdefault("labels", []); name = next((l["name"] for l in self.store.labels if l["id"] == lid), "")
        if on and lid not in ls: ls.append(lid); self.store.log(card, f"Label \"{name}\" added")
        elif not on and lid in ls: ls.remove(lid); self.store.log(card, f"Label \"{name}\" removed")
        card["updatedAt"] = now_iso(); self.store.save(); self.refresh()

    def dsub(self, sid, on):
        card, col = self._dcard()
        if not card: return
        for x in card.get("subs", []):
            if x["id"] == sid: x["done"] = on
        if card.get("subs") and all(x.get("done") for x in card["subs"]) and col != "done":
            self.store.move_card(card["id"], col, "done"); self.store.log(card, "All sub-tasks completed"); self.detail_ref = (card["id"], "done"); self.toast("🎉 All sub-tasks done — moved to Completed!")
        self.store.save(); self.refresh()

    def dsub_add(self, t):
        card, col = self._dcard(); t = t.strip()
        if not card or not t: return
        card.setdefault("subs", []).append({"id": uid(), "text": t, "done": False}); self.store.log(card, f"Checklist item added: {t}"); self.store.save(); self.refresh()

    def dsub_remove(self, sid):
        card, col = self._dcard()
        if not card: return
        card["subs"] = [x for x in card.get("subs", []) if x["id"] != sid]; self.store.save(); self.refresh()

    def dcomment(self, t):
        card, col = self._dcard(); t = t.strip()
        if not card or not t: return
        card.setdefault("comments", []).append({"id": uid(), "text": t, "at": now_iso()}); self.store.log(card, "Comment added"); self.store.save(); self.refresh()

    def dcomment_del(self, mid):
        card, col = self._dcard()
        if card: card["comments"] = [m for m in card.get("comments", []) if m["id"] != mid]; self.store.save(); self.refresh()

    # ---------- checklist ----------
    def set_cl_filter(self, k):
        self.cl_filter = k
        for kk, b in self.cl_btns.items(): b.setChecked(kk == k)
        self.render_checklist()

    def render_checklist(self):
        self._clear(self.cl_layout); s = self.store; groups = s.checklist()
        if not groups: self.cl_layout.addWidget(QLabel("No checklist groups yet. Create one to get started.", objectName="muted"))
        for g in groups:
            tasks = g["tasks"]
            if self.cl_filter == "pending": tasks = [t for t in tasks if not t.get("done")]
            if self.cl_filter == "done": tasks = [t for t in tasks if t.get("done")]
            if self.cl_q: tasks = [t for t in tasks if self.cl_q.lower() in t["title"].lower()]
            f = QFrame(); f.setObjectName("group"); fv = QVBoxLayout(f); hr = QHBoxLayout()
            tot = len(g["tasks"]); done = sum(1 for t in g["tasks"] if t.get("done")); pct = round(done / tot * 100) if tot else 0
            hr.addWidget(QLabel(f"📂  {g['name']}", objectName="sectitle")); hr.addStretch(); hr.addWidget(QLabel(f"{done}/{tot} · {pct}%", objectName="muted"))
            rn = QPushButton("✎"); rn.setObjectName("ghost"); rn.clicked.connect(lambda _, gg=g: self.rename_group(gg)); dl = QPushButton("🗑"); dl.setObjectName("ghost"); dl.clicked.connect(lambda _, gg=g: self.delete_group(gg)); hr.addWidget(rn); hr.addWidget(dl); fv.addLayout(hr)
            if not tasks: fv.addWidget(QLabel("No tasks. Add one below.", objectName="muted"))
            for t in tasks:
                r = QHBoxLayout(); cb = QCheckBox(); cb.setChecked(bool(t.get("done"))); cb.clicked.connect(lambda _, gg=g, tt=t: self.cl_toggle(gg, tt))
                body = QVBoxLayout(); tl = QLabel(t["title"]); tl.setStyleSheet("font-weight:500;" + ("text-decoration:line-through;color:#5c6382;" if t.get("done") else "")); body.addWidget(tl)
                meta = QHBoxLayout(); meta.addWidget(badge(PRI_LABEL[t.get("priority", "med")], PRI_COLOR[t.get("priority", "med")]))
                if t.get("dueDate"): od = is_overdue(t["dueDate"], t.get("done")); meta.addWidget(badge(("⚠ " if od else "📅 ") + fmt_date(t["dueDate"]), "#ff6584" if od else "#8890b0"))
                subs = t.get("subs", [])
                if subs: meta.addWidget(badge(f"{sum(1 for x in subs if x.get('done'))}/{len(subs)} subs", "#b48aff"))
                if t.get("migratedFrom"): meta.addWidget(badge("↪ " + t["migratedFrom"], "#8890b0"))
                meta.addStretch(); body.addLayout(meta)
                for x in subs:
                    scb = QCheckBox(x["text"]); scb.setChecked(bool(x.get("done"))); scb.setStyleSheet("margin-left:12px;"); scb.clicked.connect(lambda _, gg=g, tt=t, xx=x: self.cl_toggle_sub(gg, tt, xx)); body.addWidget(scb)
                eb = QPushButton("✎"); eb.setObjectName("ghost"); eb.clicked.connect(lambda _, gg=g, tt=t: self.cl_edit(gg, tt)); xb = QPushButton("✕"); xb.setObjectName("ghost"); xb.clicked.connect(lambda _, gg=g, tt=t: self.cl_delete(gg, tt))
                r.addWidget(cb, 0, Qt.AlignTop); r.addLayout(body, 1); r.addWidget(eb, 0, Qt.AlignTop); r.addWidget(xb, 0, Qt.AlignTop); fv.addLayout(r)
            ar = QHBoxLayout(); inp = QLineEdit(); inp.setPlaceholderText("+ Add task… (Enter)"); inp.returnPressed.connect(lambda gg=g, ii=inp: self.cl_add(gg, ii.text())); ab = QPushButton("Add"); ab.setObjectName("primary"); ab.clicked.connect(lambda _, gg=g, ii=inp: self.cl_add(gg, ii.text())); ar.addWidget(inp); ar.addWidget(ab); fv.addLayout(ar)
            self.cl_layout.addWidget(f)

    def add_group(self):
        name, ok = QInputDialog.getText(self, "New Checklist Group", "Group name:")
        if ok and name.strip(): self.store.checklist().append({"id": uid(), "name": name.strip(), "tasks": []}); self.store.save(); self.refresh(); self.toast("Group created!")

    def rename_group(self, g):
        name, ok = QInputDialog.getText(self, "Rename Group", "Group name:", text=g["name"])
        if ok and name.strip(): g["name"] = name.strip(); self.store.save(); self.refresh()

    def delete_group(self, g):
        if g["tasks"] and QMessageBox.question(self, "Delete group", f"Delete group \"{g['name']}\" and its {len(g['tasks'])} task(s)?") != QMessageBox.Yes: return
        self.store.checklist().remove(g); self.store.save(); self.refresh()

    def cl_add(self, g, t):
        t = t.strip()
        if t: g["tasks"].append({"id": uid(), "title": t, "done": False, "priority": "med", "dueDate": "", "subs": []}); self.store.save(); self.refresh()

    def cl_toggle(self, g, t):
        t["done"] = not t.get("done"); self.store.save(); self.refresh()
        if t["done"]: self.toast("✅ Task done!")

    def cl_toggle_sub(self, g, t, x):
        x["done"] = not x.get("done")
        if t["subs"] and all(s.get("done") for s in t["subs"]): t["done"] = True; self.toast("🎉 All sub-tasks done!")
        self.store.save(); self.refresh()

    def cl_edit(self, g, t):
        d = ChecklistTaskDialog(self, t)
        if d.exec() == QDialog.Accepted: d.apply(t); self.store.save(); self.refresh(); self.toast("Task updated")

    def cl_delete(self, g, t):
        g["tasks"].remove(t); self.store.save(); self.refresh()

    # ---------- rollover ----------
    def open_rollover(self):
        s = self.store; d = QDialog(self); d.setWindowTitle(f"Month Rollover · {s.year}"); d.setMinimumWidth(560); v = QVBoxLayout(d)
        v.addWidget(QLabel("Auto-rollover is ON: incomplete tasks from past months move forward each time the app opens. Completed tasks stay where they were finished. December → January of the next year.", wordWrap=True, objectName="muted2"))
        for mi in range(12):
            inc, incc = s.incomplete_cards(mi, s.year), s.incomplete_cl(mi, s.year); tot = len(inc) + len(incc); past = s.month_past(mi, s.year)
            ty, tm = (s.year + 1, 0) if mi == 11 else (s.year, mi + 1)
            r = QHBoxLayout(); r.addWidget(QLabel(f"{MONTHS[mi]} {s.year} → {MONTHS[tm]} {ty}", objectName="sectitle"))
            st = "🔒 Not ended yet" if not past else ("✅ All done" if tot == 0 else f"⏳ {tot} pending ({len(inc)} cards · {len(incc)} tasks)"); r.addWidget(QLabel(st, objectName="muted")); r.addStretch()
            if past and tot:
                b = QPushButton(f"Move → {MONTHS[tm]} {ty}"); b.setObjectName("primary")
                b.clicked.connect(lambda _, m=mi, dd=d: (s.migrate_month(m, s.year), self.toast("✅ Moved"), dd.accept(), self.refresh(), self.open_rollover())); r.addWidget(b)
            v.addLayout(r)
        hist = s.db["migrations"][-8:][::-1]
        if hist:
            v.addWidget(QLabel("HISTORY", objectName="navSection"))
            for m in hist: v.addWidget(QLabel(f"{'🤖' if m.get('auto') else '👆'} {MONTHS[m['from']]} {m.get('fromYear')} → {MONTHS[m['to']]} {m.get('toYear')}   +{m['kanban']} cards, +{m['checklist']} tasks   {fmt_dt(m['at'])}", objectName="muted2"))
        cb = QPushButton("Close"); cb.clicked.connect(d.accept); v.addWidget(cb); d.exec()

    # ---------- backup ----------
    def backup_now_ui(self):
        self.store.save(); dest = self.store.backup_now()
        if dest: self.toast("💾 Backup saved: " + os.path.basename(dest))

    def open_backup(self):
        d = QDialog(self); d.setWindowTitle("Backup & Restore"); d.setMinimumWidth(600); v = QVBoxLayout(d)
        v.addWidget(QLabel("Automatic every 30 min and on quit · last 30 kept in the backup folder", objectName="muted"))
        v.addWidget(QLabel("BACKUP FOLDER", objectName="navSection")); cur = get_backup_dir(); pl = QLabel(cur); pl.setWordWrap(True); pl.setObjectName("muted2"); v.addWidget(pl)
        r = QHBoxLayout()
        def choose():
            p = QFileDialog.getExistingDirectory(d, "Choose backup folder", cur)
            if p: st = read_settings(); st["backupDir"] = p; write_settings(st); d.accept(); self.toast("Backup folder: " + p); self.open_backup()
        def set_dir(p):
            st = read_settings()
            if p: st["backupDir"] = p
            else: st.pop("backupDir", None)
            write_settings(st); d.accept(); self.toast("Backups now go to " + get_backup_dir()); self.open_backup()
        cb = QPushButton("📁 Choose folder…"); cb.setObjectName("primary"); cb.clicked.connect(choose); r.addWidget(cb)
        od = detect_onedrive()
        if od: ob = QPushButton("☁ Use OneDrive"); ob.setToolTip(os.path.join(od, f"{APP_NAME} Backups")); ob.clicked.connect(lambda: set_dir(os.path.join(od, f"{APP_NAME} Backups"))); r.addWidget(ob)
        else: r.addWidget(QLabel("OneDrive not detected — use Choose folder…", objectName="muted"))
        if cur != DEFAULT_BACKUP_DIR: rb = QPushButton("Reset to default"); rb.clicked.connect(lambda: set_dir(None)); r.addWidget(rb)
        ofb = QPushButton("Open folder"); ofb.clicked.connect(lambda: self.open_path(cur)); r.addWidget(ofb); r.addStretch(); v.addLayout(r)
        lst = self.store.backup_list(); v.addWidget(QLabel(f"BACKUPS IN THIS FOLDER ({len(lst)})", objectName="navSection"))
        r2 = QHBoxLayout(); bn = QPushButton("💾 Backup now"); bn.setObjectName("primary"); bn.clicked.connect(lambda: (self.backup_now_ui(), d.accept(), self.open_backup())); r2.addWidget(bn)
        def restore_file():
            p, _ = QFileDialog.getOpenFileName(d, "Restore from backup file", cur, "Backup JSON (*.json)")
            if p: self._restore(p, d)
        rf = QPushButton("↩ Restore from file…"); rf.clicked.connect(restore_file); r2.addWidget(rf); r2.addStretch(); v.addLayout(r2)
        lw = QListWidget(); lw.setFixedHeight(200)
        for name, size, mt in lst: lw.addItem(f"{name}    {dt.datetime.fromtimestamp(mt).strftime('%Y-%m-%d %H:%M')}    {size / 1024:.1f} KB")
        v.addWidget(lw); rs = QPushButton("Restore selected"); rs.clicked.connect(lambda: self._restore(os.path.join(cur, lst[lw.currentRow()][0]), d) if lw.currentRow() >= 0 else None); v.addWidget(rs)
        v.addWidget(QLabel("LIVE DATA FILE", objectName="navSection")); v.addWidget(QLabel(DB_FILE, objectName="muted2", wordWrap=True))
        cl = QPushButton("Close"); cl.clicked.connect(d.accept); v.addWidget(cl); d.exec()

    def _restore(self, path, dlg=None):
        if QMessageBox.warning(self, "Restore backup", f"Restore from:\n{path}\n\nThis will overwrite current data (a safety backup is taken first). Continue?", QMessageBox.Yes | QMessageBox.Cancel) != QMessageBox.Yes: return
        try:
            self.store.restore_file(path)
        except Exception as e:
            QMessageBox.critical(self, "Restore failed", str(e)); return
        if dlg: dlg.accept()
        self.close_detail(); self.refresh(); self.toast("✅ Restored from " + os.path.basename(path))

    # ---------- export ----------
    def export_excel(self):
        d = ExportDialog(self, self.store)
        if d.exec() != QDialog.Accepted or not d.months(): return
        tag = MONTHS[d.months()[0]] if len(d.months()) == 1 else f"{len(d.months())}months"
        p, _ = QFileDialog.getSaveFileName(self, "Export Excel", os.path.join(os.path.expanduser("~"), "Desktop", f"{APP_NAME}_{tag}_{self.store.year}.xlsx"), "Excel (*.xlsx)")
        if not p: return
        try:
            export_excel(self.store, p, d.months(), d.inc_board.isChecked(), d.inc_cl.isChecked(), d.inc_sum.isChecked(), d.inc_done.isChecked())
        except Exception as e:
            QMessageBox.critical(self, "Export failed", str(e)); return
        self.toast("📊 Exported: " + os.path.basename(p)); self.open_path(os.path.dirname(p))

    # ---------- appearance ----------
    def palette_dict(self):
        a = self.appearance; p = dict(THEMES[a.get("theme", "dark")])
        if a.get("bg"): p["bg"] = a["bg"]
        if a.get("surface"): p["s1"] = p["s2"] = a["surface"]
        if a.get("text"): p["text"] = a["text"]
        if a.get("accent"): p["accent"] = a["accent"]
        return p

    def apply_appearance(self):
        a = self.appearance; p = self.palette_dict()
        fam = a.get("font") or ("Segoe UI" if sys.platform.startswith("win") else QApplication.font().family())
        QApplication.instance().setStyleSheet(build_qss(p, fam, int(a.get("size", 13))))
        self.theme_btn.setText("☀️" if a.get("theme") == "light" else "🌙"); self.chart.accent = p["accent"]; self.chart.update()

    def save_appearance(self):
        st = read_settings(); st["appearance"] = self.appearance; write_settings(st)
        self.store.db["meta"]["appearance_py"] = self.appearance; self.store.save(); self.apply_appearance()

    def toggle_theme(self):
        self.appearance["theme"] = "light" if self.appearance.get("theme") == "dark" else "dark"; self.save_appearance()

    def open_appearance(self):
        d = QDialog(self); d.setWindowTitle("Appearance"); d.setMinimumWidth(520); v = QVBoxLayout(d); a = self.appearance
        v.addWidget(QLabel("FONT", objectName="navSection")); fr = QHBoxLayout(); fc = QFontComboBox()
        if a.get("font"): fc.setCurrentFont(QFont(a["font"]))
        fc.currentFontChanged.connect(lambda f: (a.__setitem__("font", f.family()), self.save_appearance())); fr.addWidget(QLabel("Family")); fr.addWidget(fc, 1)
        sz = QSpinBox(); sz.setRange(10, 20); sz.setValue(int(a.get("size", 13))); sz.valueChanged.connect(lambda n: (a.__setitem__("size", n), self.save_appearance())); fr.addWidget(QLabel("Size")); fr.addWidget(sz); v.addLayout(fr)
        v.addWidget(QLabel("COLOURS", objectName="navSection"))
        def crow(label, key, varname):
            r = QHBoxLayout(); r.addWidget(QLabel(label), 1); cur = a.get(key) or self.palette_dict()[varname]; b = QPushButton(cur); b.setStyleSheet(f"background:{cur};color:#fff;")
            def pick():
                c = QColorDialog.getColor(QColor(cur), d, label)
                if c.isValid(): a[key] = c.name(); self.save_appearance(); d.accept(); self.open_appearance()
            b.clicked.connect(pick); r.addWidget(b)
            if a.get(key): rb = QPushButton("Reset"); rb.setObjectName("ghost"); rb.clicked.connect(lambda: (a.__setitem__(key, ""), self.save_appearance(), d.accept(), self.open_appearance())); r.addWidget(rb)
            v.addLayout(r)
        crow("Text colour", "text", "text"); crow("Accent colour (buttons, highlights)", "accent", "accent"); crow("Background", "bg", "bg"); crow("Cards & panels", "surface", "s2")
        tr = QHBoxLayout(); tr.addWidget(QLabel("Dark / light base theme"), 1); tb = QPushButton(("☀️ Light" if a.get("theme") == "light" else "🌙 Dark") + " — switch"); tb.clicked.connect(lambda: (self.toggle_theme(), d.accept(), self.open_appearance())); tr.addWidget(tb); v.addLayout(tr)
        v.addWidget(QLabel("PRESETS", objectName="navSection")); pg = QGridLayout()
        for i, (name, theme, vals) in enumerate(PRESETS):
            b = QPushButton(name); b.setStyleSheet(f"background:{vals.get('bg', THEMES[theme]['bg'])};color:{vals.get('text', THEMES[theme]['text'])};border-color:{vals.get('accent', '#6c8aff')};")
            b.clicked.connect(lambda _, t=theme, vv=vals: (a.update(theme=t, bg=vv.get("bg", ""), surface=vv.get("surface", ""), text=vv.get("text", ""), accent=vv.get("accent", "")), self.save_appearance(), d.accept(), self.open_appearance())); pg.addWidget(b, i // 3, i % 3)
        v.addLayout(pg)
        rr = QHBoxLayout(); ra = QPushButton("Reset everything to default"); ra.clicked.connect(lambda: (a.update(theme="dark", font="", size=13, text="", accent="", bg="", surface=""), self.save_appearance(), d.accept(), self.open_appearance())); rr.addWidget(ra); rr.addStretch(); cl = QPushButton("Close"); cl.clicked.connect(d.accept); rr.addWidget(cl); v.addLayout(rr)
        d.exec()


def main():
    app = QApplication(sys.argv); app.setApplicationName(APP_NAME); app.setQuitOnLastWindowClosed(False)
    store = Store(); w = MainWindow(store); w.show()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
